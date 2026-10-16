"""班费：统计、明细增删改、Excel 导入、发到班级群。"""
from __future__ import annotations

from typing import Any

from fastapi import APIRouter, Depends, File, Request, UploadFile
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Account
from ..security import client_context, require_admin
from ..services import class_service, fund_service, notify_service, push_service, settings_service as ss
from ..util import AppError, client_device, client_ip

router = APIRouter(prefix="/api/fund", tags=["fund"])



def fund_push_target(db: Session, class_id: int = 0) -> str:
    """班费发群的目标：该班级绑定的群（回退全局 qqbot_groups）。"""
    return ",".join(notify_service.class_targets(db, class_id=class_id))


def push_channel_text(db: Session, class_id: int = 0) -> str:
    """当前会用哪个通道发，给前端显示用。"""
    parts = []
    if ss.get_bool(db, "qqbot_enabled", False) and fund_push_target(db, class_id).strip():
        parts.append("QQ 官方机器人")
    return " + ".join(parts) if parts else "未配置"

class FundSettingsBody(BaseModel):
    per_person: float | None = None
    headcount: int | None = None
    total_fund: float | None = None
    extra_in: float | None = None
    class_name: str | None = None
    note: str | None = None
    qq_group_id: str | None = None
    template: str | None = None


class EntryBody(BaseModel):
    seq: int | None = None
    date: str | None = None
    date_text: str | None = None
    item: str = ""
    unit_price: float | None = None
    quantity: float | None = None
    amount: float = 0
    direction: str = "expense"
    handler: str = ""
    checker: str = ""
    note: str = ""


class EntryPatch(BaseModel):
    seq: int | None = None
    date: str | None = None
    date_text: str | None = None
    item: str | None = None
    unit_price: float | None = None
    quantity: float | None = None
    amount: float | None = None
    direction: str | None = None
    handler: str | None = None
    checker: str | None = None
    note: str | None = None


class DeleteBody(BaseModel):
    ids: list[int] = []


class PushBody(BaseModel):
    scope: str = "latest"  # latest / recent / selected / all
    ids: list[int] = []
    limit: int = 10
    group_id: str = ""
    title: str = ""
    dry_run: bool = False
    class_id: int = 0


# --------------------------------------------------------------------------
# 概览
# --------------------------------------------------------------------------
@router.get("/summary")
def fund_summary(request: Request, db: Session = Depends(get_db), class_id: int = 0,
                 _: Account = Depends(require_admin)):
    cid = class_service.resolve_class_id(db, class_id)
    row = fund_service.get_settings(db, cid)
    data = fund_service.summary(db, cid)
    data["class_id"] = cid
    data["template"] = row.template or fund_service.DEFAULT_TEMPLATE
    data["default_template"] = fund_service.DEFAULT_TEMPLATE
    data["group_from_settings"] = fund_push_target(db, cid)
    data["qqbot_groups"] = ss.get_setting(db, "qqbot_groups") or ""
    data["push_channel"] = push_channel_text(db, cid)
    return {"ok": True, **data}


@router.put("/settings")
def save_fund_settings(
    body: FundSettingsBody,
    request: Request,
    db: Session = Depends(get_db),
    actor: Account = Depends(require_admin),
    class_id: int = 0,
):
    cid = class_service.resolve_class_id(db, class_id)
    row = fund_service.update_settings(db, body.model_dump(exclude_none=True), cid)
    ss.write_audit(
        db,
        actor=actor.display_name or actor.username,
        role="admin",
        action="修改班费设置",
        target=row.class_name or "班费",
        detail=f"每人 {row.per_person} 元 × {row.headcount} 人",
        ip=client_ip(request),
    )
    db.commit()
    return {"ok": True, "summary": fund_service.summary(db, cid), "message": "班费设置已保存"}


# --------------------------------------------------------------------------
# 明细
# --------------------------------------------------------------------------
@router.get("/entries")
def fund_entries(
    request: Request,
    db: Session = Depends(get_db),
    keyword: str = "",
    direction: str = "",
    class_id: int = 0,
    _: Account = Depends(require_admin),
):
    cid = class_service.resolve_class_id(db, class_id)
    rows = fund_service.list_entries(db, keyword=keyword, direction=direction, class_id=cid)
    return {
        "ok": True,
        "class_id": cid,
        "items": [fund_service.serialize_entry(row) for row in rows],
        "total": len(rows),
        "summary": fund_service.summary(db, cid),
    }


@router.post("/entries")
def add_fund_entry(
    body: EntryBody,
    request: Request,
    db: Session = Depends(get_db),
    actor: Account = Depends(require_admin),
    class_id: int = 0,
):
    cid = class_service.resolve_class_id(db, class_id)
    if not body.item.strip() and not body.amount:
        raise AppError("请填写支出项目或金额")
    row = fund_service.create_entry(db, body.model_dump(), cid)
    ss.write_audit(
        db,
        actor=actor.display_name or actor.username,
        role="admin",
        action="新增班费明细",
        target=row.item,
        detail=f"{row.amount} 元 经手人 {row.handler}",
        ip=client_ip(request),
        user_agent=client_device(request),
    )
    db.commit()
    return {"ok": True, "entry": fund_service.serialize_entry(row),
            "summary": fund_service.summary(db, cid), "message": "已添加一条班费记录"}


@router.patch("/entries/{entry_id}")
def edit_fund_entry(
    entry_id: int,
    body: EntryPatch,
    request: Request,
    db: Session = Depends(get_db),
    actor: Account = Depends(require_admin),
    class_id: int = 0,
):
    cid = class_service.resolve_class_id(db, class_id)
    row = fund_service.update_entry(db, entry_id, body.model_dump(exclude_none=True))
    if row is None:
        raise AppError("这条记录不存在", 404)
    ss.write_audit(
        db,
        actor=actor.display_name or actor.username,
        role="admin",
        action="修改班费明细",
        target=row.item,
        detail=f"{row.amount} 元",
        ip=client_ip(request),
    )
    db.commit()
    return {"ok": True, "entry": fund_service.serialize_entry(row),
            "summary": fund_service.summary(db, cid), "message": "已保存"}


@router.post("/entries/delete")
def delete_fund_entries(
    body: DeleteBody,
    request: Request,
    db: Session = Depends(get_db),
    actor: Account = Depends(require_admin),
    class_id: int = 0,
):
    cid = class_service.resolve_class_id(db, class_id)
    if not body.ids:
        raise AppError("请先选择要删除的记录")
    count = fund_service.delete_entries(db, body.ids)
    ss.write_audit(
        db,
        actor=actor.display_name or actor.username,
        role="admin",
        action="删除班费明细",
        detail=f"删除 {count} 条",
        ip=client_ip(request),
    )
    db.commit()
    return {"ok": True, "deleted": count, "summary": fund_service.summary(db, cid),
            "message": f"已删除 {count} 条记录"}


# --------------------------------------------------------------------------
# 导入 / 导出
# --------------------------------------------------------------------------
@router.post("/import/preview")
async def import_preview(
    request: Request,
    file: UploadFile = File(...),
    _: Account = Depends(require_admin),
):
    data = await file.read()
    if not data:
        raise AppError("上传的文件是空的")
    try:
        parsed = fund_service.parse_workbook(data)
    except ValueError as exc:
        raise AppError(f"解析失败：{exc}") from exc
    except Exception as exc:  # noqa: BLE001
        raise AppError(f"这个文件读不了（请确认是 .xlsx 格式）：{exc}") from exc
    return {
        "ok": True,
        "filename": file.filename,
        "header_row": parsed["header_row"],
        "sheet_name": parsed["sheet_name"],
        "columns": parsed["columns"],
        "summary": parsed["summary"],
        "spent_parsed": parsed["spent_parsed"],
        "spent_in_sheet": parsed["spent_in_sheet"],
        "spent_match": parsed["spent_match"],
        "entries": parsed["entries"],
        "entry_count": len(parsed["entries"]),
    }


class ImportBody(BaseModel):
    entries: list[dict[str, Any]] = []
    summary: dict[str, Any] = {}
    replace: bool = True
    apply_summary: bool = True
    class_id: int = 0


@router.post("/import/commit")
def import_commit(
    body: ImportBody,
    request: Request,
    db: Session = Depends(get_db),
    actor: Account = Depends(require_admin),
):
    if not body.entries:
        raise AppError("没有可导入的数据，请先上传表格预览")
    cid = class_service.resolve_class_id(db, body.class_id)
    parsed = {
        "entries": body.entries,
        "summary": body.summary or {},
        "spent_parsed": round(sum(float(e.get("amount") or 0) for e in body.entries), 2),
    }
    result = fund_service.import_entries(
        db, parsed, replace=body.replace, apply_summary=body.apply_summary, class_id=cid
    )
    ss.write_audit(
        db,
        actor=actor.display_name or actor.username,
        role="admin",
        action="导入班费表格",
        detail=f"{'覆盖' if body.replace else '追加'}导入 {result['created']} 条",
        ip=client_ip(request),
        user_agent=client_device(request),
    )
    db.commit()
    return {
        "ok": True,
        "created": result["created"],
        "summary": result["summary"],
        "message": f"已导入 {result['created']} 条班费明细"
        + ("（覆盖了原有明细）" if body.replace else "（追加）"),
    }


@router.get("/export")
def export_fund(
    request: Request,
    db: Session = Depends(get_db),
    class_id: int = 0,
    _: Account = Depends(require_admin),
):
    """导出成和原表一样的班费收支明细表。"""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    from ..util import now_local

    cid = class_service.resolve_class_id(db, class_id)
    stats = fund_service.summary(db, cid)
    rows = fund_service.list_entries(db, class_id=cid)
    wb = Workbook()
    ws = wb.active
    ws.title = "班费明细"

    ws.merge_cells("A1:D2")
    ws["A1"] = "班费收支明细表"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for cell, value in (
        ("E1", "班费每人"), ("F1", "人数"), ("G1", "班费合计"), ("H1", "已使用"), ("I1", "班费余额"),
    ):
        ws[cell] = value
        ws[cell].font = Font(bold=True)
    for cell, value in (
        ("E2", stats["per_person"]), ("F2", stats["headcount"]), ("G2", stats["total_fund"]),
        ("H2", stats["spent"]), ("I2", stats["balance"]),
    ):
        ws[cell] = value

    headers = ["序号", "日期", "支出项目", "单价", "数量", "金额", "经手人", "核对人", "备注"]
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for entry in rows:
        ws.append([
            entry.seq,
            entry.date_text or (entry.entry_date.strftime("%Y.%m.%d") if entry.entry_date else ""),
            entry.item,
            entry.unit_price,
            entry.quantity,
            entry.amount,
            entry.handler,
            entry.checker,
            entry.note,
        ])
    for column, width in zip("ABCDEFGHI", (8, 14, 30, 10, 8, 12, 14, 12, 22)):
        ws.column_dimensions[column].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"班费收支明细表_{now_local().strftime('%Y%m%d')}.xlsx"
    from urllib.parse import quote

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# --------------------------------------------------------------------------
# 发到班级群
# --------------------------------------------------------------------------
@router.post("/push/preview")
def push_preview(
    body: PushBody,
    request: Request,
    db: Session = Depends(get_db),
    _: Account = Depends(require_admin),
):
    """先看看要发什么（不真的发送）。"""
    cid = class_service.resolve_class_id(db, body.class_id)
    message = fund_service.render_preview(db, body.scope, body.ids, body.limit, class_id=cid)
    row = fund_service.get_settings(db, cid)
    return {
        "ok": True,
        "text": message["text"],
        "scope_text": message["scope_text"],
        "count": message["count"],
        "amount": message["amount"],
        "entry_ids": message["entry_ids"],
        "group_id": body.group_id.strip() or fund_push_target(db, cid),
        "channel": push_channel_text(db, cid),
    }


@router.post("/push")
async def push_fund(
    body: PushBody,
    request: Request,
    db: Session = Depends(get_db),
    actor: Account = Depends(require_admin),
):
    """把班费情况通过 QQ 机器人发到班级群。"""
    cid = class_service.resolve_class_id(db, body.class_id)
    row = fund_service.get_settings(db, cid)
    message = fund_service.render_preview(db, body.scope, body.ids, body.limit, class_id=cid)
    if not message["count"]:
        raise AppError("没有可发送的明细")

    group_id = (body.group_id or fund_push_target(db, cid) or "").strip()
    if not group_id:
        raise AppError(
            "还没有可用的推送通道：请到「系统配置 → QQ 官方机器人」填好群 openid（在群里 @ 机器人就能拿到）"
        )

    if body.dry_run:
        return {"ok": True, "text": message["text"], "sent": None, "message": "仅预览，没有发送"}

    sent = await push_service.deliver_notification(
        db,
        title=body.title.strip() or "班费公示",
        content=message["text"],
        scene="fund_report",
        group_openid=group_id,
    )

    if sent.get("ok"):
        fund_service.stamp_pushed(db, message["entry_ids"])

    # 记进推送日志，方便在「提醒推送 → 推送日志」里查看
    try:
        notify_service.log_push(
            db,
            channel=sent.get("channel") or "none",
            scene="fund_report",
            target=group_id,
            title=body.title.strip() or "班费公示",
            content=message["text"],
            status="success" if sent.get("ok") else "failed",
            detail=sent.get("message", ""),
            trigger="manual",
            class_id=cid,
        )
    except Exception:  # noqa: BLE001
        pass
    ss.write_audit(
        db,
        actor=actor.display_name or actor.username,
        role="admin",
        action="班费发群",
        target=group_id or "（邮件通道）",
        detail=f"{message['scope_text']}，{message['amount']} 元",
        ip=client_ip(request),
        user_agent=client_device(request),
    )
    db.commit()
    return {
        "ok": bool(sent.get("ok")),
        "text": message["text"],
        "scope_text": message["scope_text"],
        "count": message["count"],
        "amount": message["amount"],
        "group_id": group_id,
        "channel": sent.get("channel"),
        "result": sent,
        "message": "已发送到班级群" if sent.get("ok") else sent.get("message", "发送失败"),
    }
