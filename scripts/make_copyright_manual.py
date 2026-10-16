#!/usr/bin/env python3
"""生成软著申请用《操作手册》。"""
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"E:\Document\code\宿舍管理系统")
OUT_DIR = ROOT / "软著申请"
SHOTS = OUT_DIR / "screenshots"
OUT = OUT_DIR / "操作手册-生活助手宿舍管理系统V1.0.docx"

SOFTWARE = "生活助手宿舍管理系统"
VERSION = "V1.0"
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
DARK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x5A, 0x63, 0x72)

USABLE_CM = 16.6


def set_font(run, size=10.5, bold=None, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), 15, True, ACCENT)
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), 12.5, True, DARK)
    return p


def para(text, size=10.5, color=None, bold=False, indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    set_font(p.add_run(text), size, bold, color)
    return p


def bullet(text, head=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if head:
        set_font(p.add_run(head), 10.5, True)
    set_font(p.add_run(text), 10.5)
    return p


def step(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(f"{num}. "), 10.5, True, ACCENT)
    set_font(p.add_run(text), 10.5)
    return p


def table(rows, widths=None, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run(str(val)), 9.5, header and i == 0)
    if widths:
        total = sum(widths) or 1.0
        scaled = [w / total * USABLE_CM for w in widths]
        for row in t.rows:
            for idx, w in enumerate(scaled):
                row.cells[idx].width = Cm(w)
    return t


def picture(name, caption, width_cm=15.6):
    path = SHOTS / f"{name}.png"
    if not path.exists():
        para(f"[缺少截图：{name}]", 9.5, GREY)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)
    set_font(cp.add_run(caption), 9, False, GREY)


def add_page_number(paragraph):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText")
    i.set(qn("xml:space"), "preserve")
    i.text = "PAGE"
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run._r.append(b)
    run._r.append(i)
    run._r.append(e)
    set_font(run, 9, False, GREY)


# ==================== 文档初始化 ====================
doc = Document()
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.line_spacing = 1.42
style.paragraph_format.space_after = Pt(4)

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(hp.add_run(f"{SOFTWARE} {VERSION}    操作手册"), 9, False, GREY)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("第 "), 9, False, GREY)
add_page_number(fp)
set_font(fp.add_run(" 页"), 9, False, GREY)

# ==================== 封面 ====================
for _ in range(4):
    doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(t.add_run(SOFTWARE), 26, True, ACCENT)
v = doc.add_paragraph()
v.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(v.add_run(VERSION), 18, True, DARK)
s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.paragraph_format.space_before = Pt(6)
set_font(s.add_run("操 作 手 册"), 20, True, DARK)

for _ in range(3):
    doc.add_paragraph()

info = doc.add_table(rows=0, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, val in (
    ("软件名称", SOFTWARE),
    ("版本号", VERSION),
    ("软件分类", "应用软件 — 管理信息系统"),
    ("开发语言", "Python 3.12、JavaScript (ES6)、HTML5、CSS3"),
    ("开发框架", "FastAPI、SQLAlchemy、Uvicorn、APScheduler"),
    ("运行平台", "Linux（Ubuntu 24.04）/ Windows，跨平台"),
    ("文档类型", "操作手册（用户手册）"),
):
    cells = info.add_row().cells
    for idx, txt in enumerate((k, val)):
        cells[idx].text = ""
        p = cells[idx].paragraphs[0]
        set_font(p.add_run(txt), 10.5, idx == 0, DARK if idx == 0 else None)
    cells[0].width = Cm(3.6)
    cells[1].width = Cm(11.0)

doc.add_page_break()

# ==================== 目录 ====================
h1("目　录")
toc = [
    "一、软件概述",
    "二、运行环境与安装部署",
    "三、用户角色与访问入口",
    "四、公开端功能操作说明",
    "五、管理端功能操作说明",
    "六、定时任务与自动处理",
    "七、数据存储与备份",
    "八、常见问题处理",
    "九、功能清单",
    "十、错误提示与处理",
    "十一、主要页面字段说明",
    "十二、维护与升级",
    "十三、安全设计说明",
    "附录 A　数据库表结构",
    "附录 B　主要接口清单",
]
for item in toc:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.6)
    set_font(p.add_run(item), 11)

# ==================== 一、软件概述 ====================
doc.add_page_break()
h1("一、软件概述")

h2("1.1 编制目的")
para("本手册是「生活助手宿舍管理系统 V1.0」的配套操作文档，用于说明该软件的功能构成、"
     "安装部署方式与各项操作流程，供使用者、维护人员及软件著作权登记审查参考。")

h2("1.2 软件简介")
para("生活助手宿舍管理系统是一套面向高校班级宿舍日常管理的 Web 应用系统。"
     "系统以「宿舍」为最小管理单元，为班级生活委员、班主任及全体同学提供"
     "每日在寝情况填报、自动汇总统计、异常情况记录、班费收支管理与消息推送等功能。")
para("系统采用 B/S 架构，用户通过浏览器访问即可使用，无需安装客户端。"
     "系统区分公开端与管理端两个独立访问入口，分别面向学生与管理者的使用场景。")

h2("1.3 主要功能")
table(
    [
        ["功能模块", "功能说明"],
        ["人员信息管理", "以宿舍为组导入、新增、编辑、删除人员信息（姓名、学号、电话、宿舍号）"],
        ["每日填报", "寝室长或同学勾选在寝情况、填写缺席原因与备注，提交当日宿舍情况"],
        ["自动统计", "按日、周、月统计填报率、在寝人数、缺席人次，并生成可视化图表"],
        ["群内汇报", "每晚定时将当日填报汇总自动推送到班级 QQ 群，便于班主任查看"],
        ["未填提醒", "定时检测未填报宿舍，向宿舍长推送提醒消息"],
        ["班费管理", "导入班费收支明细表，支持增删改查、汇总公示与一键发群"],
        ["Excel 导入导出", "支持名单模板下载、批量导入名单、导出台账与填报记录"],
        ["账号与权限", "管理员账号与宿舍长账号分离，宿舍可设置独立验证密码"],
        ["消息推送", "支持 QQ 官方机器人（QQ 开放平台）推送群汇报与定时提醒"],
        ["数据备份", "定时自动备份数据库，支持手动下载备份文件"],
        ["操作日志", "记录登录、填报、配置变更等关键操作，便于追溯"],
    ],
    widths=[3.6, 13.0],
)

h2("1.4 技术架构")
table(
    [
        ["层次", "技术选型"],
        ["前端", "原生 HTML5 / CSS3 / JavaScript (ES6)，单页应用 + 哈希路由，Chart.js 图表"],
        ["后端", "Python 3.12 + FastAPI（异步框架）+ SQLAlchemy ORM"],
        ["数据库", "SQLite（WAL 模式），单文件部署，支持事务与并发读"],
        ["服务容器", "Uvicorn（ASGI 服务器），单进程双端口监听"],
        ["定时任务", "APScheduler（AsyncIOScheduler）"],
        ["机器人接入", "QQ 官方机器人开放平台（WebSocket 长连接 + HTTP OpenAPI）"],
        ["报表处理", "openpyxl（Excel 读写）、segno（二维码生成）"],
        ["安全机制", "itsdangerous 会话令牌、PBKDF2 密码散列（20 万次迭代）、PyNaCl 签名校验"],
    ],
    widths=[2.6, 14.0],
)

h2("1.5 软件规模")
data = [
    ["类别", "文件数", "代码行数"],
    ["Python 源程序", "57", "13649"],
    ["JavaScript 源程序", "17", "4263"],
    ["CSS 样式", "1", "362"],
    ["HTML 页面", "1", "102"],
    ["合计", "76", "18376"],
]
table(data, widths=[5.0, 4.0, 4.0])
para("说明：以上为项目全部源程序统计；其中核心业务与前端代码 42 个文件、"
     "约 13219 行，作为软件著作权登记的源程序材料提交。", 9.5, GREY)

# ==================== 二、运行环境与安装 ====================
doc.add_page_break()
h1("二、运行环境与安装部署")

h2("2.1 运行环境要求")
table(
    [
        ["项目", "要求"],
        ["操作系统", "Linux（推荐 Ubuntu 20.04 及以上）或 Windows 10/11"],
        ["运行环境", "Python 3.10 及以上（推荐 3.12）"],
        ["内存", "建议 2 GB 及以上"],
        ["磁盘", "建议 5 GB 以上可用空间"],
        ["浏览器", "Chrome / Edge / Firefox / Safari 等现代浏览器（支持手机浏览器）"],
        ["网络", "服务器需可访问外网（用于 QQ 官方机器人推送）"],
    ],
    widths=[3.4, 13.2],
)

h2("2.2 安装步骤")
step(1, "安装 Python 3.12，并创建虚拟环境：python -m venv .venv")
step(2, "安装依赖：pip install -r requirements.txt")
step(3, "配置环境变量（可选）：DORM_HOST、DORM_PORT、DORM_ADMIN_PORT、DORM_DATA_DIR、DORM_TIMEZONE")
step(4, "初始化数据库：首次启动时系统自动建表并写入默认配置")
step(5, "启动服务：python run.py（同时监听公开端口与管理端口）")
step(6, "浏览器访问：公开端 http://服务器地址:8090，管理端 http://服务器地址:8091")

h2("2.3 主要环境变量")
table(
    [
        ["变量名", "默认值", "说明"],
        ["DORM_HOST", "0.0.0.0", "监听地址"],
        ["DORM_PORT", "8090", "公开端端口（同学填报）"],
        ["DORM_ADMIN_PORT", "8091", "管理端端口（管理员）"],
        ["DORM_DATA_DIR", "./data", "数据目录（数据库、备份、导出、日志）"],
        ["DORM_TIMEZONE", "Asia/Shanghai", "时区"],
        ["DORM_APP_NAME", "生活助手", "系统显示名称"],
    ],
    widths=[3.8, 3.0, 9.8],
)

h2("2.4 双端口设计说明")
para("系统通过同一进程监听两个端口，实现公开端与管理端的物理隔离：")
bullet("公开端（默认 8090）：面向全体同学与宿舍长，提供填报、看板与查询功能；")
bullet("管理端（默认 8091）：面向生活委员与班主任，提供全部管理与配置功能；")
bullet("管理接口在公开端一律返回 404，管理端未登录访问一律返回 401，双重隔离。")

# ==================== 三、角色与入口 ====================
doc.add_page_break()
h1("三、用户角色与访问入口")

h2("3.1 角色说明")
table(
    [
        ["角色", "访问入口", "权限范围"],
        ["普通同学", "公开端", "浏览宿舍列表、查看填报进度、代为填报未设密码的宿舍"],
        ["宿舍长", "公开端", "本宿舍填报与修改；宿舍若设密码则仅本宿舍长可写"],
        ["管理员（生活委员/班主任）", "管理端", "全部数据查看、编辑、导出、配置与推送"],
    ],
    widths=[4.4, 3.0, 9.2],
)

h2("3.2 登录与身份识别")
para("普通同学访问公开端无需注册登录，点击自己所在宿舍即可填报。")
para("宿舍长账号的用户名即宿舍号（如 6-213），密码由管理员设置，也可不设置。")
para("管理员使用独立账号登录管理端。若使用宿舍长账号尝试登录管理端，"
     "系统将拒绝并提示前往公开端登录，避免越权访问。")

# ==================== 四、公开端操作 ====================
doc.add_page_break()
h1("四、公开端功能操作说明")

h2("4.1 登录页")
para("公开端首页顶部提供两个入口：直接进入填报，或使用宿舍长账号登录。")
picture("01-登录页", "图 4-1　登录页")

h2("4.2 选择宿舍")
para("页面列出全部宿舍及其宿舍长、人数与位置信息。可输入宿舍号快速查找，"
     "点击宿舍卡片进入填报页面。")
picture("02-选择寝室", "图 4-2　选择宿舍页面")

h2("4.3 每日填报")
para("填报页面分为左侧填报区与右侧进度区。")
step(1, "选择填报日期（默认当天，可补填最近 7 天）")
step(2, "逐个确认宿舍成员是否在寝，点击「在寝 / 不在」切换")
step(3, "对不在寝的人员选择原因（回家、病假、事假、外出、实习、其他）")
step(4, "在备注栏填写不在人员说明或意外情况")
step(5, "点击「提交今晚填报」完成提交")
para("系统对同一宿舍同一日期只保留一条记录，重复提交会覆盖原有内容，"
     "不会产生重复数据。如需从头填写，可点击「全部重置为「在寝」」。", 10)
picture("03-每日填报", "图 4-3　每日填报页面")

h2("4.4 数据看板")
para("看板展示当日填报进度：已填报宿舍数、未填报宿舍数、在寝人数、缺席人次，"
     "并列出尚未填报的宿舍名单。")
picture("04-数据看板", "图 4-4　数据看板")

h2("4.5 统计图表")
para("支持按日、周、月三个维度查看历史数据趋势，包括填报率变化、"
     "缺席情况分布、各宿舍表现对比等，以柱状图、折线图、饼图等形式呈现。")
picture("05-统计图表", "图 4-5　统计图表")

h2("4.6 填报记录")
para("按日期查看各宿舍填报详情，支持一键导出 Excel 台账，"
     "便于长期留档与上报。")
picture("06-填报记录", "图 4-6　填报记录")

# ==================== 五、管理端操作 ====================
doc.add_page_break()
h1("五、管理端功能操作说明")

h2("5.1 管理员登录")
para("访问管理端地址，输入管理员账号与密码登录。系统对管理端全部接口实施统一鉴权，"
     "未登录状态下任何数据接口均返回 401，前端登录页无法被绕过。")
picture("07-管理端登录", "图 5-1　管理端登录页")

h2("5.2 数据总览")
para("登录后进入数据总览，展示宿舍数、学生数、今日填报情况、班费概况等关键指标。")
picture("08-管理端总览", "图 5-2　管理端数据总览")

h2("5.3 宿舍与人员管理")
para("以宿舍为单位管理全部人员信息，支持：")
bullet("新增、编辑、删除宿舍；")
bullet("新增、编辑、删除人员（姓名、学号、电话、宿舍号、职务、备注）；")
bullet("下载名单导入模板、批量导入名单、导出名单；")
bullet("设置宿舍长姓名与联系电话。")
picture("09-宿舍与人员", "图 5-3　宿舍与人员管理")

h2("5.4 账号管理")
para("为每间宿舍的宿舍长创建登录账号（用户名即宿舍号），可设置、重置、"
     "清空密码，也可停用账号。密码留空表示不设密码，此时任何人都可代为填报该宿舍。")
picture("10-账号管理", "图 5-4　账号管理")

h2("5.5 班费管理")
para("班费模块以导入的 Excel 收支明细表为数据基准，支持：")
bullet("导入班费收支明细表并自动识别表头与汇总数据；")
bullet("新增、编辑、删除费用明细；")
bullet("查看班费总额、已使用金额、当前余额与人均金额；")
bullet("将最新一笔、最近若干笔、选定若干笔或全部班费情况一键发送到班级群；")
bullet("导出台账 Excel。")
picture("11-班费管理", "图 5-5　班费管理")

h2("5.6 提醒推送")
para("展示每晚自动汇报与未填提醒的实际内容，支持手动补发；"
     "页面下方记录每次推送的渠道、目标、结果与错误原因。")
picture("12-提醒推送", "图 5-6　提醒推送")

h2("5.7 管理控制台")
para("集中展示全部宿舍的宿舍长密码状态，支持批量设置密码、重置密码、"
     "批量创建宿舍长账号。")
picture("13-管理控制台", "图 5-7　管理控制台")

h2("5.8 填写日志")
para("记录每一次填报与操作的时间、宿舍、操作者、来源设备与 IP，便于核查与追溯。")
picture("14-填写日志", "图 5-8　填写日志")

h2("5.9 系统配置")
para("系统全部参数均可在网页上图形化修改，无需修改代码。分为基本设置、"
     "QQ 官方机器人、消息模板、定时任务、"
     "宿舍长联系表、系统与备份等页签。")
picture("15-系统配置-基本", "图 5-9　系统配置 — 基本设置")

para("QQ 官方机器人配置页提供长连接状态、凭证检测、群列表获取与测试消息发送，"
     "并显示最近上报的群 openid，便于一键填入推送目标。")
picture("16-系统配置-QQ机器人", "图 5-10　系统配置 — QQ 官方机器人")

# ==================== 六、定时任务 ====================
doc.add_page_break()
h1("六、定时任务与自动处理")
para("系统内置定时任务调度，全部时间均可在「系统配置 → 定时任务」中修改：")
table(
    [
        ["任务", "默认时间", "功能说明"],
        ["未填报提醒", "22:00", "检测未填报宿舍并推送提醒"],
        ["每晚群汇报", "22:40", "汇总当日填报情况并推送到班级 QQ 群"],
        ["自动导出 Excel", "23:30", "自动导出当日填报台账并存档"],
        ["数据库备份", "03:00", "自动备份数据库，默认保留最近 30 份"],
    ],
    widths=[3.4, 2.4, 10.8],
)
para("消息推送统一走 QQ 官方机器人（QQ 开放平台），无需挂 QQ 客户端、不会被封号。"
     "配额不足或未获授权时接口会返回错误码，程序会把失败原因记入推送日志。", 10)
para("系统对网络抖动具备容错能力：发送失败自动重试，长连接断开自动重连，"
     "会话令牌失效自动刷新。", 10)

# ==================== 七、数据与备份 ====================
h1("七、数据存储与备份")

h2("7.1 数据存储")
table(
    [
        ["目录/文件", "内容"],
        ["app.db", "SQLite 主数据库文件（WAL 模式）"],
        ["secret.key", "会话签名密钥"],
        ["backups/", "数据库备份文件"],
        ["exports/", "自动导出的 Excel 台账"],
        ["uploads/", "导入的名单与明细表原件"],
        ["logs/", "服务运行日志"],
    ],
    widths=[3.6, 13.0],
)

h2("7.2 备份机制")
para("系统通过定时任务每日自动备份数据库，采用 SQLite 官方备份接口，"
     "先将 WAL 日志合并进主库再备份，确保备份内容完整。"
     "备份文件按时间命名，超出保留份数的旧备份自动清理。")
para("管理员亦可在「系统与备份」页面手动触发备份并下载备份文件。")

h2("7.3 数据完整性设计")
bullet("填报表对「宿舍 + 日期」建立唯一约束，从数据库层面杜绝重复记录；")
bullet("高频查询字段均建立索引，全部热点查询走索引，无全表扫描；")
bullet("数据库启用 WAL 日志模式与写锁等待，支持多人同时操作；")
bullet("并发提交冲突时使用保存点机制自动降级为更新，不会产生脏数据。")

# ==================== 八、常见问题 ====================
doc.add_page_break()
h1("八、常见问题处理")
qa = [
    ("同学打不开网页", "确认访问地址与端口正确；若网络受限，可改用备用地址访问。"),
    ("找不到自己的宿舍", "联系管理员在「宿舍与人员」中核实并补充宿舍信息。"),
    ("填错了想重填", "同一天重新提交会覆盖原记录；也可点击「全部重置为「在寝」」后重新填写。"),
    ("群里没收到自动汇报", "在「提醒推送」页面点击「立即发送到 QQ 群」手动补发，并查看下方失败原因。"),
    ("宿舍长忘记密码", "管理员在「账号管理」或「管理控制台」中重置密码。"),
    ("管理员忘记密码", "在服务器上执行 scripts/admin_account.py --reset 重置。"),
    ("数据会不会丢失", "数据保存在服务器，且每日自动备份；管理员也可随时手动下载备份。"),
    ("能不能用手机填报", "可以，系统为响应式设计，手机浏览器可直接访问使用。"),
]
for q, a in qa:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    set_font(p.add_run("问　" + q), 10.5, True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(0.5)
    p2.paragraph_format.space_after = Pt(2)
    set_font(p2.add_run("答　" + a), 10.5)

# ==================== 九、功能清单 ====================
doc.add_page_break()
h1("九、功能清单")
para("下表列出系统的全部功能点，可作为软件功能完整性的核对依据。")
table(
    [
        ["编号", "功能点", "所属模块", "说明"],
        ["F-01", "宿舍信息维护", "宿舍与人员", "新增、编辑、删除宿舍，设置宿舍长与联系电话"],
        ["F-02", "人员信息维护", "宿舍与人员", "新增、编辑、删除人员，含姓名/学号/电话/宿舍号/职务/备注"],
        ["F-03", "名单模板下载", "宿舍与人员", "下载标准 Excel 导入模板"],
        ["F-04", "名单批量导入", "宿舍与人员", "上传 Excel 自动识别列并导入宿舍与人员"],
        ["F-05", "名单导出", "宿舍与人员", "导出全部宿舍与人员信息为 Excel"],
        ["F-06", "填写日志查询", "日志", "按时间查看填报与操作记录，含设备与 IP"],
        ["F-07", "今日进度查看", "数据看板", "展示已填/未填宿舍、在寝人数、缺席人次"],
        ["F-08", "未填宿舍名单", "数据看板", "实时列出尚未填报的宿舍"],
        ["F-09", "日维度统计", "统计图表", "按日统计填报率与人数变化"],
        ["F-10", "周维度统计", "统计图表", "按周汇总填报情况"],
        ["F-11", "月维度统计", "统计图表", "按月汇总填报情况"],
        ["F-12", "图表可视化", "统计图表", "柱状图、折线图、饼图展示趋势与对比"],
        ["F-13", "在线填报", "填报", "选择宿舍、逐人勾选在寝情况"],
        ["F-14", "缺席原因登记", "填报", "选择原因类型并填写备注"],
        ["F-15", "补填历史", "填报", "可补填最近 7 天的填报记录"],
        ["F-16", "填报覆盖", "填报", "同宿舍同日期重复提交自动覆盖"],
        ["F-17", "一键重置", "填报", "一键将所有成员重置为在寝状态"],
        ["F-18", "填报记录查询", "填报记录", "按日期范围与宿舍筛选查看"],
        ["F-19", "填报记录导出", "填报记录", "导出 Excel 台账"],
        ["F-20", "填报记录删除", "填报记录", "删除指定填报记录（管理员）"],
        ["F-21", "QQ 群自动汇报", "推送", "定时将当日汇总发送到班级群"],
        ["F-22", "未填提醒推送", "推送", "定时检测未填宿舍并推送提醒"],
        ["F-23", "手动补发", "推送", "页面手动触发一次汇报或提醒"],
        ["F-24", "推送日志", "推送", "记录每次推送的渠道、目标、结果与错误"],
        ["F-25", "机器人长连接", "推送", "WebSocket 长连接收发群消息，断线自动重连"],
        ["F-26", "群内 @ 指令", "推送", "群内 @ 机器人可查询汇报或记录填报"],
        ["F-27", "定时提醒", "推送", "可配置多条定时提醒，每条自定义时间与内容"],
        ["F-28", "班费明细导入", "班费", "导入 Excel 收支明细表并识别汇总"],
        ["F-29", "班费明细维护", "班费", "新增、编辑、删除费用条目"],
        ["F-30", "班费汇总统计", "班费", "统计总额、已用、余额、人均"],
        ["F-31", "班费公示发群", "班费", "按范围一键发送到班级群"],
        ["F-32", "班费导出", "班费", "导出班费台账 Excel"],
        ["F-33", "宿舍长账号管理", "账号", "创建、停用、重置账号与密码"],
        ["F-34", "宿舍验证密码", "权限", "设置/重置/清空宿舍验证密码"],
        ["F-35", "密码批量生成", "账号", "为未设密码宿舍批量生成随机密码"],
        ["F-36", "管理员密码修改", "账号", "管理员自行修改登录密码"],
        ["F-37", "参数图形化配置", "系统配置", "全部参数网页化修改，无需改代码"],
        ["F-38", "定时任务配置", "系统配置", "修改各定时任务的执行时间"],
        ["F-39", "消息模板配置", "系统配置", "自定义汇报与提醒文案模板"],
        ["F-40", "手动数据库备份", "系统备份", "一键备份并下载数据库文件"],
        ["F-41", "自动数据库备份", "定时任务", "每日定时备份，超出份数自动清理"],
        ["F-42", "数据导出归档", "定时任务", "每日自动导出 Excel 台账"],
        ["F-43", "数据清空", "系统备份", "清空业务数据（需二次确认）"],
        ["F-44", "双端口隔离", "安全", "公开端与管理端端口级物理隔离"],
        ["F-45", "接口统一鉴权", "安全", "管理端全部接口强制身份校验"],
        ["F-46", "操作审计", "安全", "记录关键操作的操作者、时间、IP"],
    ],
    widths=[1.6, 3.4, 2.8, 8.8],
)

# ==================== 十、错误提示与处理 ====================
doc.add_page_break()
h1("十、错误提示与处理")
table(
    [
        ["提示信息", "原因", "处理方式"],
        ["这个宿舍还没有人员，无法填报", "该宿舍尚未导入人员", "先在「宿舍与人员」中导入或添加人员"],
        ["有人不在时，请在备注里写清原因", "存在不在寝人员但备注为空", "补充缺席原因后重新提交"],
        ["不能填报未来的日期", "选择的日期晚于当天", "改为当天或历史日期"],
        ["只能补填最近 7 天", "补填日期超出 7 天", "在 7 天范围内选择日期"],
        ["该宿舍已由宿舍长设置密码，填报需要先登录", "该宿舍已设验证密码", "使用宿舍长账号登录后再填写"],
        ["请先用管理员账号登录管理端", "管理端未登录或会话过期", "重新登录管理端"],
        ["这个账号不是管理员，没有管理权限", "使用宿舍长账号访问管理端", "改用管理员账号，或到公开端登录"],
        ["这是寝室长账号，不能登录管理端", "宿舍长账号尝试登录管理端", "到公开端登录"],
        ["还没有填写 AppID / AppSecret", "未配置 QQ 机器人凭证", "在「系统配置 → QQ 官方机器人」填写"],
        ["发送超时（网络到腾讯不稳）", "服务器到腾讯网络抖动", "系统已自动重试；仍失败则检查服务器网络"],
        ["该群没有开通机器人主动消息权限", "群未授权机器人主动发言", "让群主在群内授权，或改用群内 @ 机器人"],
        ["连接 bots.qq.com 超时", "外网访问受限", "检查服务器网络与安全组出站策略"],
        ["密码至少 6 位", "设置的密码过短", "改为 6 位及以上密码"],
        ["宿舍号已存在", "重复添加同一宿舍", "改为其他宿舍号或直接编辑现有宿舍"],
        ["导入失败：表头识别不到", "Excel 表头不规范", "使用系统提供的名单模板重新整理"],
    ],
    widths=[5.4, 4.2, 7.0],
)

# ==================== 十一、页面字段说明 ====================
doc.add_page_break()
h1("十一、主要页面字段说明")

h2("11.1 填报页面")
table(
    [
        ["字段", "类型", "说明"],
        ["填报日期", "日期选择", "默认当天，可选最近 7 天"],
        ["宿舍", "下拉选择", "当前填报的宿舍"],
        ["填表人", "文本", "可选填，用于记录提交人"],
        ["在寝 / 不在", "按钮切换", "逐人标记在寝状态"],
        ["缺席原因", "下拉选择", "回家、病假、事假、外出、实习、其他"],
        ["备注", "多行文本", "说明不在人员或意外情况"],
    ],
    widths=[3.2, 2.8, 10.6],
)

h2("11.2 宿舍与人员页面")
table(
    [
        ["字段", "必填", "说明"],
        ["宿舍号", "是", "唯一标识，如 6-213"],
        ["楼栋 / 位置", "否", "如 新区校区6号学生公寓213"],
        ["宿舍长", "否", "宿舍长姓名"],
        ["联系电话", "否", "宿舍长手机号"],
        ["姓名", "是", "学生姓名"],
        ["学号", "否", "学生学号"],
        ["电话", "否", "学生手机号"],
        ["职务", "否", "成员 / 宿舍长"],
        ["备注", "否", "其他说明"],
    ],
    widths=[3.2, 1.8, 11.6],
)

h2("11.3 班费页面")
table(
    [
        ["字段", "说明"],
        ["日期", "费用发生日期"],
        ["摘要", "费用用途说明"],
        ["金额", "正数为支出，负数为收入"],
        ["经手人", "经办人姓名"],
        ["备注", "补充说明"],
        ["是否已推送", "标记该笔是否已发群公示"],
    ],
    widths=[3.6, 13.0],
)

# ==================== 十二、维护与升级 ====================
doc.add_page_break()
h1("十二、维护与升级")

h2("12.1 日常维护")
bullet("通过「数据看板」每日确认填报情况；")
bullet("定期在「系统与备份」页面下载数据库备份留档；")
bullet("查看服务器日志目录，关注推送失败记录；")
bullet("定期检查磁盘空间，必要时清理历史导出文件。")

h2("12.2 服务管理命令")
table(
    [
        ["操作", "命令"],
        ["启动服务", "systemctl start dorm"],
        ["停止服务", "systemctl stop dorm"],
        ["重启服务", "systemctl restart dorm"],
        ["查看状态", "systemctl status dorm"],
        ["查看实时日志", "journalctl -u dorm -f"],
        ["查看应用日志", "tail -f /www/dorm-data/logs/service.log"],
    ],
    widths=[3.4, 13.2],
)

h2("12.3 数据备份与恢复")
para("备份：系统每日 03:00 自动备份；也可在管理端「系统与备份」页面手动触发并下载。")
para("恢复：停止服务后，用备份文件覆盖数据目录下的 app.db，再启动服务即可。")
para("建议每月将备份文件下载到本地或其他存储介质，做异地留存。")

h2("12.4 版本升级")
step(1, "先备份数据库与当前程序目录；")
step(2, "替换程序文件（保留 data 数据目录不动）；")
step(3, "如依赖有变化，重新执行 pip install -r requirements.txt；")
step(4, "重启服务：systemctl restart dorm；")
step(5, "验证：访问公开端与管理端，确认功能正常。")
para("系统启动时会自动检测并补齐新增的数据库字段，升级不需要手动改表结构。", 10)

# ==================== 十三、安全设计说明 ====================
h1("十三、安全设计说明")
table(
    [
        ["安全措施", "说明"],
        ["双端口物理隔离", "公开端与管理端由同一进程监听不同端口，管理接口在公开端一律返回 404"],
        ["接口统一鉴权", "管理端除登录等少数接口外，全部接口强制校验管理员会话，未登录返回 401"],
        ["角色权限校验", "写操作校验账号角色，宿舍长账号无法执行管理操作"],
        ["密码散列存储", "密码采用 PBKDF2 算法加盐散列，20 万次迭代，不保存明文"],
        ["会话令牌", "采用 itsdangerous 签名令牌，有效期 30 天，支持手动失效"],
        ["密钥不外泄", "公开端配置接口仅返回白名单字段，AppSecret、邮箱密码等一律不返回"],
        ["并发一致性", "唯一约束 + 事务保存点，保证多人同时操作不产生重复数据"],
        ["操作审计", "登录、填报、配置变更等关键操作均记录操作者、时间、设备与 IP"],
        ["签名校验", "机器人事件回调采用 Ed25519 签名校验，防止伪造请求"],
        ["数据备份", "每日自动备份数据库，防止数据意外丢失"],
    ],
    widths=[3.8, 12.8],
)

# ==================== 附录 A ====================
doc.add_page_break()
h1("附录 A　数据库表结构")
table(
    [
        ["表名", "说明", "主要字段"],
        ["dorms", "宿舍表", "宿舍号、楼栋、宿舍长、联系电话、验证密码、排序、启用状态"],
        ["members", "人员表", "宿舍ID、姓名、学号、电话、宿舍号、职务、备注、启用状态"],
        ["reports", "填报记录表", "宿舍ID、日期、是否齐全、缺席人员、备注、人数统计、来源、提交人、IP、设备"],
        ["report_checkins", "填报明细表", "记录ID、人员ID、姓名、是否在寝、原因"],
        ["accounts", "账号表", "用户名、显示名、密码散列、角色、绑定宿舍、电话、启用状态、最近登录信息"],
        ["fund_entries", "班费明细表", "序号、日期、摘要、金额、经手人、备注、是否已推送"],
        ["fund_settings", "班费设置表", "人均金额、人数、备注说明"],
        ["bot_groups", "机器人群表", "渠道、群 openid、群号、群名称、最近消息、消息数"],
        ["settings", "系统配置表", "配置键、值、分组、说明"],
        ["push_logs", "推送日志表", "渠道、场景、目标、标题、内容、状态、详情、触发方式"],
        ["audit_logs", "操作日志表", "操作者、角色、动作、对象、详情、IP、设备、时间"],
    ],
    widths=[3.0, 2.4, 11.2],
)

# ==================== 附录 B ====================
h1("附录 B　主要接口清单")
para("系统共提供 REST 接口 70 余个，按模块划分如下：")
table(
    [
        ["模块", "主要接口"],
        ["认证与账号", "/api/auth/login、/api/auth/logout、/api/auth/me、/api/auth/password、/api/accounts"],
        ["访问状态", "/api/access/status、/api/access/leader/login、/api/access/leader/password"],
        ["宿舍与人员", "/api/dorms、/api/members、/api/members/import、/api/members/export"],
        ["填报", "/api/reports/submit、/api/reports/today、/api/reports/dorm/{id}、/api/reports/records、/api/reports/stats、/api/reports/export"],
        ["班费", "/api/fund/summary、/api/fund/entries、/api/fund/import、/api/fund/export、/api/fund/push"],
        ["机器人", "/api/qqbot/events、/api/qqbot/status、/api/qqbot/groups"],
        ["系统", "/api/system/info、/api/system/mode、/api/system/backup、/api/system/reset-data、/api/settings、/api/jobs、/api/logs/push"],
    ],
    widths=[3.0, 13.6],
)

end = doc.add_paragraph()
end.paragraph_format.space_before = Pt(20)
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(end.add_run(f"— {SOFTWARE} {VERSION} 操作手册 完 —"), 10, False, GREY)

doc.save(OUT)
print("已生成:", OUT)
print("大小:", f"{os.path.getsize(OUT) / 1024:.1f} KB")
