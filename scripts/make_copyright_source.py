#!/usr/bin/env python3
"""生成软著申请用《源程序》文档：
   - A4，每页 50 行，页眉含软件名称+版本号，页脚连续页码
   - 提交前 30 页 + 后 30 页（共 60 页）
"""
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"E:\Document\code\宿舍管理系统")
OUT_DIR = ROOT / "软著申请"
OUT_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "源程序-生活助手宿舍管理系统V1.0.docx"

SOFTWARE = "生活助手宿舍管理系统"
VERSION = "V1.0"
LINES_PER_PAGE = 50
FIRST_PAGES = 30
LAST_PAGES = 30

# 收录的核心源代码（按逻辑顺序：入口 → 配置 → 模型 → 安全 → 路由 → 服务 → 前端）
SOURCE_FILES = [
    "run.py",
    "backend/config.py",
    "backend/database.py",
    "backend/models.py",
    "backend/defaults.py",
    "backend/util.py",
    "backend/security.py",
    "backend/main.py",
    "backend/scheduler.py",
    "backend/routers/access.py",
    "backend/routers/auth.py",
    "backend/routers/classes.py",
    "backend/routers/dorms.py",
    "backend/routers/reports.py",
    "backend/routers/fund.py",
    "backend/routers/qqbot.py",
    "backend/routers/settings.py",
    "backend/services/report_service.py",
    "backend/services/class_service.py",
    "backend/services/fund_service.py",
    "backend/services/notify_service.py",
    "backend/services/push_service.py",
    "backend/services/qqbot_gateway.py",
    "backend/services/settings_service.py",
    "backend/services/export_service.py",
    "backend/services/backup_service.py",
    "backend/excel_util.py",
    "backend/static/index.html",
    "backend/static/css/style.css",
    "backend/static/js/core.js",
    "backend/static/js/app.js",
    "backend/static/js/pages/home.js",
    "backend/static/js/pages/dashboard.js",
    "backend/static/js/pages/report.js",
    "backend/static/js/pages/overview.js",
    "backend/static/js/pages/records.js",
    "backend/static/js/pages/stats.js",
    "backend/static/js/pages/fund.js",
    "backend/static/js/pages/accounts.js",
    "backend/static/js/pages/dorms.js",
    "backend/static/js/pages/classes.js",
    "backend/static/js/pages/console.js",
    "backend/static/js/pages/push.js",
    "backend/static/js/pages/logs.js",
    "backend/static/js/pages/config.js",
    "backend/static/js/pages/login.js",
]


MAX_UNITS = 100  # 一行最多 100 个半角宽度（中文算 2）


def display_width(text: str) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)


def truncate_display(text: str, limit: int) -> str:
    if display_width(text) <= limit:
        return text
    out = []
    used = 0
    for ch in text:
        w = 2 if ord(ch) > 0x2E80 else 1
        if used + w > limit - 1:
            break
        out.append(ch)
        used += w
    return "".join(out) + "…"


def collect_lines() -> tuple[list[str], int]:
    """把源码收集成统一的文本行，并在文件之间加分隔注释。"""
    lines: list[str] = []
    total = 0
    for rel in SOURCE_FILES:
        path = ROOT / rel
        if not path.exists():
            print(f"  ! 缺少文件：{rel}")
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        body = text.splitlines()
        total += len(body)
        lines.append(f"// ===== 文件：{rel} =====")
        for ln in body:
            # 按显示宽度截断：中文等全角字符算 2 个宽度，
            # 保证一行在 A4（8pt Consolas）里不会自动换行，
            # 否则一页放不下 50 行，不符合软著要求。
            lines.append(truncate_display(ln, MAX_UNITS))
        lines.append("")
    return lines, total


def set_cell_or_run_font(run, name="Consolas", size=8):
    run.font.name = name
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_page_number(paragraph):
    """在段落里插入 PAGE 域。"""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_cell_or_run_font(run, "宋体", 9)


def build_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run(f"{SOFTWARE} {VERSION}    源程序")
    set_cell_or_run_font(run, "宋体", 9)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("第 ")
    set_cell_or_run_font(r1, "宋体", 9)
    add_page_number(fp)
    r2 = fp.add_run(" 页")
    set_cell_or_run_font(r2, "宋体", 9)


def main():
    print("收集源码 ...")
    lines, total = collect_lines()
    print(f"  收录文件 {len(SOURCE_FILES)} 个，源码约 {total} 行，展开后 {len(lines)} 行")

    # 按每页 50 行切分
    pages: list[list[str]] = []
    for i in range(0, len(lines), LINES_PER_PAGE):
        pages.append(lines[i : i + LINES_PER_PAGE])
    total_pages = len(pages)
    print(f"  折合 {total_pages} 页（每页 {LINES_PER_PAGE} 行）")

    if total_pages <= FIRST_PAGES + LAST_PAGES:
        picked = list(range(total_pages))
        print(f"  总页数不足 60，提交全部 {total_pages} 页")
    else:
        # 软著规范：源程序前、后各连续 30 页（不足 60 页提交全部），
        # 每页不少于 50 行 ——「结束页除外」，所以末页是残页是允许的。
        picked = list(range(FIRST_PAGES)) + list(range(total_pages - LAST_PAGES, total_pages))
        print(f"  提交前 {FIRST_PAGES} 页 + 后 {LAST_PAGES} 页 = {len(picked)} 页")
        print(f"    （末页为程序结尾残页 {len(pages[picked[-1]])} 行，规范允许）")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(8)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(13.6)  # 固定行高，保证一页放得下 50 行

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(1.6)
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.8)
    build_header_footer(sec)

    for idx, page_no in enumerate(picked):
        for line in pages[page_no]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(13.6)
            run = p.add_run(line if line.strip() else " ")
            set_cell_or_run_font(run, "Consolas", 8)
        if idx < len(picked) - 1:
            brk = doc.add_paragraph()
            brk.paragraph_format.space_before = Pt(0)
            brk.paragraph_format.space_after = Pt(0)
            brk.paragraph_format.line_spacing = Pt(1)
            brun = brk.add_run()
            set_cell_or_run_font(brun, "Consolas", 1)
            brun.add_break(WD_BREAK.PAGE)

    doc.save(OUT)
    print(f"\n已生成：{OUT}")
    print(f"大小：{os.path.getsize(OUT) / 1024:.1f} KB")
    return OUT


if __name__ == "__main__":
    main()
