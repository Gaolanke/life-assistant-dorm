"""端到端自检脚本：不需要浏览器，直接打接口验证全部功能。

认证方式：账号登录（管理员 admin + 密码），管理接口带 `X-Session-Token`。
脚本开头会把 admin 的密码重置成一个已知值，所以可以反复运行。

用法： .venv\\Scripts\\python.exe scripts\\selftest.py [http://127.0.0.1:8091]

注意：最后一节会清空业务数据（宿舍 / 人员 / 填报），跑完请重新导入名单：
    .venv\\Scripts\\python.exe scripts\\import_roster.py "新区宿舍物联网25-2(1).xlsx" --yes
    .venv\\Scripts\\python.exe scripts\\make_history.py --days 14 --fresh
"""
from __future__ import annotations

import io
import sys
from datetime import date, timedelta
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# 自动对齐服务的数据目录（服务的 DORM_DATA_DIR 写在 systemd 单元里）
sys.path.insert(0, str(Path(__file__).resolve().parent))
import server_env  # noqa: E402

if str(Path(__file__).resolve().parent.parent) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

print(f"数据目录：{server_env.DATA_DIR}")

from auth_client import (  # noqa: E402
    ADMIN_BASE,
    PUBLIC_BASE,
    TEST_ADMIN_PASSWORD,
    ensure_admin_client,
    public_client,
    try_login,
)

BASE = sys.argv[1] if len(sys.argv) > 1 else ADMIN_BASE
PUBLIC = sys.argv[2] if len(sys.argv) > 2 else PUBLIC_BASE
if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass

PASSED: list[str] = []
FAILED: list[str] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    if condition:
        PASSED.append(name)
        print(f"  [OK]   {name}")
    else:
        FAILED.append(f"{name} :: {detail}")
        print(f"  [FAIL] {name}  {detail}")


def make_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "学号", "电话号", "宿舍号"])
    ws.append(["张三", "2023010101", "13800000001", "3-201"])
    ws.append(["李四", "2023010102", "13800000002", "3-201"])
    ws.append(["王五", "2023010103", "13800000003", "3-201"])
    ws.append(["赵六", "2023010104", "13800000004", "3-202"])
    ws.append(["钱七", "2023010105", "13800000005", "3-202"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def main() -> int:
    print("\n== 0. 确认打的是管理端口 ==")
    probe = httpx.Client(base_url=BASE, timeout=30, follow_redirects=True)
    try:
        mode = probe.get("/api/system/mode").json()
    except Exception as exc:  # noqa: BLE001
        print(f"  连不上 {BASE}：{exc}")
        probe.close()
        return 2
    print(
        f"  当前地址 {BASE} → {mode.get('message')}"
        f"（公开端口 {mode.get('main_port')} / 管理端口 {mode.get('admin_port')}）"
    )
    if not mode.get("is_admin_port"):
        admin_base = BASE.replace(str(mode.get("main_port")), str(mode.get("admin_port")))
        print(
            "\n  ✗ 这是公开端口，管理接口只在管理端口开放。\n"
            f"    请改成： python scripts/selftest.py {admin_base}\n"
        )
        probe.close()
        return 2
    probe.close()

    print("\n== 1. 基础、账号登录与权限 ==")
    guest = httpx.Client(base_url=BASE, timeout=30, follow_redirects=True)
    health = guest.get("/api/health").json()
    check("健康检查", health.get("ok") is True, str(health))

    access = guest.get("/api/access/status").json()
    check("未登录时不是管理员", access["is_admin"] is False, str(access))
    check("未登录时 logged_in=False", access["logged_in"] is False, str(access))
    check("接口提示必须登录", access.get("login_required") is True, str(access))

    no_token = guest.get("/api/accounts")
    check("未登录访问管理接口返回 401", no_token.status_code == 401, f"状态码 {no_token.status_code}")

    code, data = try_login(base=BASE, username="admin", password="definitely-wrong")
    check("错误密码登录被拒绝", code == 403, f"状态码 {code}")
    code, data = try_login(base=BASE, username="no-such-user", password="whatever")
    check("不存在的账号登录被拒绝", code == 403, f"状态码 {code}")

    client = ensure_admin_client(base=BASE)
    _token = client.headers.get("X-Session-Token", "")
    check("管理员登录成功", bool(_token), "登录接口没有返回 token")

    me = client.get("/api/auth/me").json()
    check("登录后 /api/auth/me 认得我", me.get("logged_in") is True, str(me))
    check("当前账号是管理员", (me.get("account") or {}).get("role") == "admin", str(me.get("account")))
    check("登录后能看到账号列表", client.get("/api/accounts").status_code == 200, "账号列表不可用")

    guest_me = guest.get("/api/auth/me").json()
    check("没带令牌时 /api/auth/me 显示未登录", guest_me.get("logged_in") is False, str(guest_me))

    with httpx.Client(
        base_url=BASE, timeout=30, headers={"X-Session-Token": "fake-token-123"}
    ) as fake:
        check("伪造令牌被拒绝", fake.get("/api/accounts").status_code == 401, "伪造令牌竟然通过了")

    open_write = client.post("/api/dorms", json={"room_no": "9-999"})
    check("登录后可以新增寝室", open_write.status_code == 200 and open_write.json()["ok"], open_write.text[:200])
    if open_write.status_code == 200:
        client.delete(f"/api/dorms/{open_write.json()['item']['id']}?purge_reports=true")

    with httpx.Client(base_url=BASE, timeout=30, follow_redirects=True) as temp:
        temp.post("/api/auth/login", json={"username": "admin", "password": TEST_ADMIN_PASSWORD})
        logout = temp.post("/api/auth/logout")
        check("退出登录成功", logout.status_code == 200 and logout.json()["ok"], logout.text[:200])
        check(
            "退出后不再是登录状态",
            temp.get("/api/auth/me").json().get("logged_in") is False,
            "退出后依然登录着",
        )

    print("\n== 2. 模板下载与导入 ==")
    template = client.get("/api/members/template")
    check("模板下载成功", template.status_code == 200 and len(template.content) > 2000, str(template.status_code))

    upload = client.post(
        "/api/members/import/preview",
        files={"file": ("member_test.xlsx", make_workbook(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    preview = upload.json()
    check("导入预览成功", upload.status_code == 200 and preview.get("ok"), upload.text[:200])
    check("识别出 2 间宿舍", preview.get("total_dorms") == 2, str(preview.get("total_dorms")))
    check("识别出 5 人", preview.get("total_members") == 5, str(preview.get("total_members")))
    check("列映射正确", preview["mapping"].get("name") == "姓名" and preview["mapping"].get("room_no") == "宿舍号", str(preview.get("mapping")))

    commit = client.post("/api/members/import/commit", json={"token": preview["token"], "strategy": "merge", "update_member": True})
    check("导入提交成功", commit.status_code == 200 and commit.json()["ok"], commit.text[:200])
    check("新增 2 间宿舍", commit.json()["created_dorms"] == 2, str(commit.json()))

    print("\n== 3. 宿舍与人员 CRUD ==")
    dorms = client.get("/api/dorms").json()
    rooms = [d["room_no"] for d in dorms["items"]]
    check("宿舍列表含导入的宿舍", "3-201" in rooms and "3-202" in rooms, str(rooms))
    dorm = next(d for d in dorms["items"] if d["room_no"] == "3-201")
    check("宿舍成员数正确", dorm["member_count"] == 3, str(dorm["member_count"]))
    check("宿舍长字段存在", "leader_name" in dorm, str(dorm.get("leader_name")))

    new_dorm = client.post("/api/dorms", json={"room_no": "3-203", "building": "3号楼", "leader_name": "孙八", "leader_phone": "13900000008"})
    check("新增宿舍成功", new_dorm.status_code == 200 and new_dorm.json()["ok"], new_dorm.text[:200])
    dorm203 = new_dorm.json()["item"]

    dup = client.post("/api/dorms", json={"room_no": "3-203"})
    check("重复宿舍号被拒绝", dup.status_code == 400, str(dup.status_code))

    member = client.post(f"/api/dorms/{dorm203['id']}/members", json={"name": "周九", "student_no": "2023010109", "phone": "13900000009", "role": "成员"})
    check("新增成员成功", member.status_code == 200 and member.json()["ok"], member.text[:200])
    member_id = member.json()["item"]["id"]

    edited = client.put(f"/api/members/{member_id}", json={"name": "周九", "student_no": "2023010109", "phone": "13900000099", "room_no": "3-203", "role": "宿舍长", "active": True})
    check("编辑成员成功", edited.status_code == 200 and edited.json()["item"]["phone"] == "13900000099", edited.text[:200])

    renamed = client.put(f"/api/dorms/{dorm203['id']}", json={"room_no": "3-204", "building": "3号楼", "leader_name": "孙八", "leader_phone": "13900000008", "qq_group_id": "123456789", "active": True})
    check("修改宿舍号成功", renamed.status_code == 200 and renamed.json()["item"]["room_no"] == "3-204", renamed.text[:200])
    dorm204 = renamed.json()["item"]

    members_204 = client.get("/api/dorms").json()["items"]
    dorm_204 = next(d for d in members_204 if d["room_no"] == "3-204")
    leader_member = next((m for m in dorm_204["members"] if m["name"] == "周九"), None)
    if leader_member:
        leader_res = client.post(f"/api/dorms/{dorm_204['id']}/leader?member_id={leader_member['id']}", json={})
        check("可以指定宿舍长（无需账号）", leader_res.status_code == 200 and leader_res.json()["ok"], leader_res.text[:200])
        check("宿舍长同步到宿舍信息", leader_res.json()["item"]["leader_name"] == "周九", leader_res.text[:200])
    else:
        check("可以指定宿舍长（无需账号）", False, "没找到成员 周九")
        check("宿舍长同步到宿舍信息", False, "没找到成员 周九")

    print("\n== 4. 时间同步 ==")
    info = client.get("/api/system/info").json()
    check("系统信息返回服务器时间", bool(info.get("server_time")), str(info)[:120])
    check("系统信息返回时区", bool(info.get("timezone")), str(info)[:120])
    today_res = client.get("/api/reports/today").json()
    check("填报接口也返回服务器时间", bool(today_res.get("server_time")), today_res.get("server_time", ""))

    print("\n== 5. 填报表单 ==")
    form = client.get(f"/api/reports/dorm/{dorm['id']}").json()
    check("填报页能取到成员", len(form["members"]) == 3, str(len(form["members"])))
    check("缺席原因候选项存在", len(form["absence_reasons"]) >= 3, str(form["absence_reasons"]))

    checkins = [{"member_id": m["id"], "name": m["name"], "present": m["name"] != "李四", "reason": "回家" if m["name"] == "李四" else ""} for m in form["members"]]
    submit = client.post("/api/reports/submit", json={
        "dorm_id": dorm["id"],
        "report_date": date.today().isoformat(),
        "is_complete": False,
        "absent_names": "李四",
        "note": "李四回家（家长已知晓）",
        "checkins": checkins,
        "submitter": "生活委员",
    })
    check("提交填报成功", submit.status_code == 200 and submit.json()["ok"], submit.text[:300])
    report = submit.json()["report"]
    check("缺席人数正确", report["absent_count"] == 1, str(report))
    check("在寝人数正确", report["present_count"] == 2, str(report))
    check("记录带提交人", report["submitter"] == "生活委员", str(report["submitter"]))

    resubmit = client.post("/api/reports/submit", json={
        "dorm_id": dorm["id"], "report_date": date.today().isoformat(), "is_complete": True,
        "absent_names": "", "note": "已回寝", "checkins": [],
    })
    check("重复提交会覆盖", resubmit.status_code == 200 and resubmit.json()["report"]["is_complete"] is True, resubmit.text[:200])

    today = client.get("/api/reports/today").json()
    check("今日统计包含填报", today["stats"]["reported"] >= 1, str(today["stats"]))
    check("能看到未填报宿舍", isinstance(today["stats"]["missing_rooms"], list), str(today["stats"]))

    dash = client.get("/api/reports/dashboard").json()
    check("看板返回各宿舍当天明细", len(dash.get("items") or []) >= 2, str(len(dash.get("items") or [])))
    check("看板明细含填报状态", all("status_text" in row for row in dash["items"]), str(dash["items"][:1]))

    print("\n== 6. 无鉴权：任何人都能填报任意宿舍（公开端口） ==")
    outsider = public_client(PUBLIC)
    other = next(d for d in client.get("/api/dorms").json()["items"] if d["room_no"] == "3-201")
    outside = outsider.post("/api/reports/submit", json={
        "dorm_id": other["id"], "report_date": date.today().isoformat(), "is_complete": True,
        "note": "无", "submitter": "路过的同学",
    })
    check("可以填报任意宿舍（不需要登录）", outside.status_code == 200 and outside.json()["ok"], outside.text[:200])
    check("填表人已记录", outside.json()["report"]["submitter"] == "路过的同学", outside.text[:200])

    own_detail = client.get(f"/api/reports/dorm/{dorm204['id']}").json()
    own_submit = outsider.post("/api/reports/submit", json={
        "dorm_id": dorm204["id"], "report_date": date.today().isoformat(), "is_complete": True, "note": "无",
        "checkins": [{"member_id": m["id"], "present": True} for m in own_detail["members"]],
        "submitter": "",
    })
    check("不留名字也能提交（记为匿名）", own_submit.status_code == 200 and own_submit.json()["report"]["submitter"] == "匿名填报", own_submit.text[:200])

    header_submit = outsider.post(
        "/api/reports/submit",
        headers={"X-Client-Name": "%E6%9D%A8%E6%9C%9D%E8%BE%89"},  # “杨朝辉”的 URL 编码
        json={"dorm_id": other["id"], "report_date": date.today().isoformat(), "is_complete": True, "note": "无"},
    )
    check("请求头里的填表人也能识别（URL 编码）", header_submit.status_code == 200 and header_submit.json()["report"]["submitter"] == "杨朝辉", header_submit.text[:200])

    print("\n== 7. 统计与导出 ==")
    stats = client.get("/api/reports/stats?granularity=day&points=7").json()
    check("按天统计可用", len(stats["items"]) == 7, str(len(stats.get("items", []))))
    check("周统计可用", len(client.get("/api/reports/stats?granularity=week&points=4").json()["items"]) == 4)
    check("月统计可用", len(client.get("/api/reports/stats?granularity=month&points=3").json()["items"]) == 3)
    check("排行榜有数据", len(stats["ranking_week"]) >= 2, str(len(stats.get("ranking_week", []))))

    export = client.get("/api/reports/export")
    check("导出今日 Excel", export.status_code == 200 and export.content[:2] == b"PK", str(export.status_code))
    wb = load_workbook(io.BytesIO(export.content))
    sheets = wb.sheetnames
    check("导出含汇总表", any("汇总" in s or s == date.today().isoformat() for s in sheets), str(sheets))
    check("导出含明细表", any("明细" in s for s in sheets), str(sheets))
    summary = wb[sheets[0]]
    header_row = [c.value for c in summary[4]] if summary.max_row >= 4 else []
    check("汇总表列头正确", "宿舍号" in header_row and "是否齐全" in header_row, str(header_row))

    rng = client.get(f"/api/reports/export?start={(date.today() - timedelta(days=2)).isoformat()}&end={date.today().isoformat()}")
    check("导出区间 Excel", rng.status_code == 200 and rng.content[:2] == b"PK", str(rng.status_code))
    zip_res = client.get(f"/api/reports/export?start={(date.today() - timedelta(days=2)).isoformat()}&end={date.today().isoformat()}&as_zip=true")
    check("导出 zip 打包", zip_res.status_code == 200 and zip_res.content[:2] == b"PK", str(zip_res.status_code))

    now_export = client.post("/api/reports/export-now")
    check("手动导出到服务器成功", now_export.status_code == 200 and now_export.json()["ok"], now_export.text[:200])
    files = client.get("/api/reports/export-files").json()
    check("服务器文件列表有内容", len(files["items"]) >= 1, str(files)[:200])
    dl = client.get(f"/api/reports/export-files/{files['items'][0]['name']}")
    check("下载服务器文件成功", dl.status_code == 200 and dl.content[:2] == b"PK", str(dl.status_code))

    print("\n== 8. 记录查询 ==")
    records = client.get("/api/reports/records?start=" + (date.today() - timedelta(days=3)).isoformat()).json()
    check("记录查询可用", records["total"] >= 2, str(records["total"]))
    check("缺勤明细可用", isinstance(records["absents"], list), str(type(records["absents"])))

    print("\n== 9. 推送与模板（模拟模式） ==")
    settings = client.get("/api/settings").json()
    check("配置项已初始化", "qq_report_time" in settings["values"], str(list(settings["values"])[:5]))
    check("模板变量说明存在", "dorm_lines" in settings["template_vars"], str(len(settings.get("template_vars", {}))))

    saved = client.put("/api/settings", json={"values": {"qq_report_time": "22:40", "school_name": "测试大学"}, "reschedule": True})
    check("修改配置生效", saved.json()["values"]["school_name"] == "测试大学", saved.text[:200])

    jobs = client.get("/api/jobs").json()
    ids = [j["id"] for j in jobs["jobs"]]
    # 定时任务：群汇报 + 自动导出 + 数据库备份（外加「提醒推送」里配置的 reminder_* 任务）
    # 旧的「填报提醒」(fill_remind)、「未填报提醒」(missing_reminder) 已删除，这里断言它们不在
    check(
        "定时任务已注册",
        all(k in ids for k in ("group_report", "auto_export", "db_backup")),
        str(ids),
    )
    check("旧的填报提醒任务已删除", "fill_remind" not in ids, str(ids))
    check("未填报提醒任务已删除", "missing_reminder" not in ids, str(ids))
    report_job = next(j for j in jobs["jobs"] if j["id"] == "group_report")
    check("汇报任务下次执行时间是 22:40", report_job["next_run"].endswith("22:40:00"), str(report_job))

    preview = client.get("/api/reports/preview").json()
    check("群汇报预览非空", len(preview["group_report"]) > 20, preview["group_report"][:100])
    check("提醒预览非空", len(preview["reminder"]) > 10, preview["reminder"][:100])
    check("预览含学校名", "测试大学" in preview["group_report"], preview["group_report"][:200])

    # QQ 官方机器人测试接口：只校验接口结构正确、有明确说明文案（不发真实消息）
    bad_bot = client.post("/api/settings/test", json={"kind": "qqbot"})
    bot_body = bad_bot.json()
    check(
        "机器人测试接口返回结构正确",
        bad_bot.status_code == 200 and "ok" in bot_body and bool(bot_body.get("message")),
        bad_bot.text[:200],
    )

    logs = client.get("/api/logs/push").json()
    check("推送日志已记录失败原因", len(logs["items"]) >= 2, str(len(logs["items"])))
    check("日志里有排错信息", any(item["detail"] for item in logs["items"]), str(logs["items"])[:200])

    print("\n== 10. 账号体系（新的登录方式） ==")
    accounts = client.get("/api/accounts").json()
    check("账号列表可用", accounts["ok"] and accounts["total"] >= 1, str(accounts)[:200])
    check("账号列表里有 admin", any(item["username"] == "admin" for item in accounts["items"]), str(accounts["items"])[:200])
    check("账号列表标出超级管理员", any(item["is_owner"] for item in accounts["items"]), str(accounts["items"])[:200])
    check("账号统计可用", client.get("/api/accounts/stats").status_code == 200, "账号统计不可用")
    check("未登录看不到账号列表", httpx.get(f"{BASE}/api/accounts", timeout=20).status_code == 401, "未登录竟然能看账号")
    check("公开端口没有账号接口", public_client(PUBLIC).get("/api/accounts").status_code == 404, "公开端口不应该有 /api/accounts")
    check("旧接口 /api/access/code 已下线", client.get("/api/access/code").status_code == 404, str(client.get("/api/access/code").status_code))
    check("旧接口 /api/access/devices 已下线", client.get("/api/access/devices").status_code == 404, str(client.get("/api/access/devices").status_code))
    check("旧接口 /api/access/verify 已下线", client.get("/api/access/verify").status_code == 404, str(client.get("/api/access/verify").status_code))
    check("填写日志接口可用", client.get("/api/logs/submissions").status_code == 200, "填写日志不可用")
    check("操作日志新地址可用", client.get("/api/logs/audit").status_code == 200, "操作日志不可用")
    check("操作日志老地址兼容", client.get("/api/audit").status_code == 200, "老地址不可用")
    check("班费概览可用", client.get("/api/fund/summary").status_code == 200, "班费概览不可用")

    print("\n== 12. 系统信息与备份 ==")
    info = client.get("/api/system/info").json()
    check("系统信息可用", info["ok"] and info["counts"]["dorms"] >= 3, str(info["counts"]))
    check("系统信息里带账号数", info["counts"].get("accounts", 0) >= 1, str(info["counts"]))
    backup = client.get("/api/system/backup")
    check("数据库备份可下载", backup.status_code == 200 and backup.content[:6] == b"SQLite" and len(backup.content) > 4096, f"{len(backup.content)} {backup.content[:16]!r}")
    audit = client.get("/api/audit").json()
    check("操作日志有记录", len(audit["items"]) >= 5, str(len(audit["items"])))
    check("日志记录了来源 IP", any(item["ip"] for item in audit["items"]), "no ip")
    qr = client.get("/api/system/qrcode?text=" + BASE)
    check("二维码生成可用", qr.status_code == 200 and qr.content[:4] == b"\x89PNG", str(qr.status_code))

    print("\n== 13. 静态页面 ==")
    index = client.get("/")
    check("首页可访问", index.status_code == 200 and "生活助手" in index.text, str(index.status_code))
    for asset, keyword in (("/css/style.css", "--blue"), ("/js/core.js", "App.api"), ("/js/app.js", "App.init"), ("/vendor/chart.umd.min.js", "Chart")):
        res = client.get(asset)
        check(f"静态资源 {asset}", res.status_code == 200 and keyword in res.text, str(res.status_code))
    check("登录页面资源存在", client.get("/js/pages/login.js").status_code == 200, "login.js 缺失")

    print("\n== 14. 清空业务数据 ==")
    reset = client.post("/api/system/reset-data", json={"keep_accounts": True, "keep_settings": True})
    check("清空业务数据成功", reset.status_code == 200 and reset.json()["ok"], reset.text[:200])
    check("清空后没有宿舍", client.get("/api/dorms").json()["total"] == 0, "dorms not cleared")
    check("清空后配置保留", "qq_report_time" in client.get("/api/settings").json()["values"], "settings lost")
    check("清空后账号保留", client.get("/api/accounts").json()["total"] >= 1, "accounts lost")
    check("清空后仍能登录", client.get("/api/auth/me").json().get("logged_in") is True, "login lost")

    client.close()
    guest.close()
    outsider.close()

    print("\n" + "=" * 62)
    print(f"通过 {len(PASSED)} 项，失败 {len(FAILED)} 项")
    if FAILED:
        print("\n失败明细：")
        for item in FAILED:
            print("  - " + item)
        return 1
    print("全部通过 ✅")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
