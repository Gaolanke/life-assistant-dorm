#!/usr/bin/env python3
"""生成软著申请《申请表填写信息》和《材料清单与办理指引》。"""
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(r"E:\Document\code\宿舍管理系统\软著申请")
SOFTWARE = "生活助手宿舍管理系统"
VERSION = "V1.0"
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
DARK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x5A, 0x63, 0x72)


def new_doc(title: str):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.line_spacing = 1.4
    style.paragraph_format.space_after = Pt(4)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(f"{SOFTWARE} {VERSION}    {title}")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(9)
    r.font.color.rgb = GREY

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(i); run._r.append(e)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    return doc


def title1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = ACCENT


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = ACCENT


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = DARK


def para(doc, text, size=10.5, color=None, bold=False, indent=0.0):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = color
    return p


def bullet(doc, text, head=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if head:
        r = p.add_run(head)
        r.font.name = "微软雅黑"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        r.font.size = Pt(10.5)
        r.font.bold = True
    r2 = p.add_run(text)
    r2.font.name = "微软雅黑"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r2.font.size = Pt(10.5)
    return p


def table(doc, rows, widths):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name = "微软雅黑"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            r.font.size = Pt(9.5)
            r.font.bold = i == 0
    total = sum(widths)
    scaled = [w / total * 16.6 for w in widths]
    for row in t.rows:
        for idx, w in enumerate(scaled):
            row.cells[idx].width = Cm(w)
    return t


# ==========================================================
# 文档一：申请表填写信息
# ==========================================================
doc = new_doc("申请表填写信息")
title1(doc, "软件著作权登记申请表 · 填写信息")
para(doc, "以下内容为登录「中国版权保护中心」软件著作权登记系统后，"
          "需要在申请表中逐项填写的信息。可直接照抄。", 10, GREY)
para(doc, "")

h1(doc, "一、软件基本信息")
table(doc, [
    ["申请表项目", "填写内容"],
    ["软件全称", SOFTWARE],
    ["软件简称", "生活助手（可不填）"],
    ["版本号", VERSION],
    ["软件分类", "应用软件"],
    ["软件作品说明", "原创"],
    ["开发完成日期", "2026 年 10 月 11 日"],
    ["首次发表日期", "2026 年 10 月 11 日"],
    ["发表状态", "已发表"],
    ["发表地点", "中国 甘肃省 兰州市（通过网络发表）"],
    ["开发方式", "独立开发"],
    ["权利取得方式", "原始取得"],
    ["权利范围", "全部权利"],
], [5.0, 11.0])

h1(doc, "二、软件功能与技术特点")
para(doc, "软件用途和技术特点（约 100–300 字，可直接填写）：", 10, GREY)
para(doc,
     "本软件是一套面向高校班级宿舍日常管理的 Web 应用系统，以宿舍为最小管理单元，"
     "服务于班级生活委员、班主任及全体学生。系统采用 B/S 架构与前后端分离设计，"
     "由同一进程分别监听公开端与管理端两个端口，实现管理权限的物理隔离。"
     "主要功能包括：以宿舍为单位的人员信息导入与维护；学生每日在寝情况在线填报，"
     "支持逐人勾选、缺席原因与备注记录；按日、周、月自动汇总统计并生成可视化图表；"
     "每晚定时将填报汇总推送至班级 QQ 群，便于班主任掌握情况；对未填报宿舍定时提醒；"
     "班费收支明细的导入、编辑、公示与一键发群；填报台账与人员名单的 Excel 导入导出；"
     "宿舍长账号与验证密码管理；数据库定时自动备份与操作日志追溯。"
     "系统采用 FastAPI 异步框架与 SQLAlchemy ORM，数据库使用 SQLite（WAL 模式），"
     "对关键查询字段建立索引，并通过唯一约束与事务保存点机制保证多人并发操作下的数据一致性。")

h1(doc, "三、软件开发信息（在线表「软件开发信息」页，每项限 50 字）")
para(doc, "以下内容已按 50 字上限精确核算，可直接复制粘贴到在线表单对应输入框。", 10, GREY)
table(doc, [
    ["表单字段（必填）", "填写内容", "字数"],
    ["开发的硬件环境", "8核CPU/40GB内存/1024GB硬盘的PC", "24"],
    ["运行的硬件环境", "2核CPU/2GB内存/40GB硬盘服务器", "21"],
    ["开发该软件的操作系统", "Windows 11 64位", "14"],
    ["软件开发环境/开发工具", "VS Code、Git、Python 3.12", "23"],
    ["该软件的运行平台/操作系统", "Windows、Linux（Ubuntu 24.04）", "27"],
    ["软件运行支撑环境/支持软件", "Python 3.12、SQLite 3、Chrome等浏览器", "31"],
], [4.6, 9.4, 1.6])

h1(doc, "四、软件功能与特点（在线表「软件功能与特点」页）")
para(doc, "该页字数上限较宽，以下内容可直接使用。", 10, GREY)

h2(doc, "软件用途")
para(doc,
     "用于高校班级宿舍日常管理：宿舍人员信息维护、学生每日在寝情况在线填报、"
     "按日周月自动统计汇总、定时向班级群推送填报结果、班费收支管理、"
     "名单与台账 Excel 导入导出。")

h2(doc, "技术特点")
para(doc,
     "采用B/S架构与前后端分离设计；同一进程分别监听公开端与管理端两个端口，"
     "实现权限物理隔离；后端基于FastAPI异步框架与SQLAlchemy ORM，"
     "数据库采用SQLite（WAL模式）；对高频查询字段建立索引，"
     "并通过唯一约束与事务保存点机制保证并发操作下的数据一致性；"
     "内置定时任务调度与多通道消息推送，发送失败自动降级重试；"
     "支持数据库定时自动备份与操作日志追溯；前端为响应式设计，手机浏览器可直接使用。")

h2(doc, "开发目的")
para(doc,
     "替代传统的逐间询问、手工统计方式，实现宿舍情况填报、汇总、提醒与公示的自动化，"
     "减轻班级生活委员与班主任的管理负担，提高数据准确性与时效性。")

h2(doc, "面向领域/行业")
para(doc, "教育 — 学校班级管理")

h1(doc, "五、著作权人信息")
para(doc, "以下信息需按实际情况填写（个人申请填自然人信息，学校/单位申请填法人信息）。", 10, GREY)
table(doc, [
    ["申请表项目", "填写内容"],
    ["著作权人类型", "□ 自然人　　□ 法人或其他组织　（按实际情况勾选）"],
    ["姓名或名称", "（填写你的真实姓名或学校名称）"],
    ["证件类型", "居民身份证 / 统一社会信用代码证书"],
    ["证件号码", "（填写对应证件号码）"],
    ["国籍/地区", "中国"],
    ["省份/城市", "甘肃省 兰州市"],
    ["联系地址", "（填写可接收快递的地址）"],
    ["邮政编码", "（填写）"],
    ["联系人 / 电话 / 邮箱", "（填写）"],
], [4.6, 11.4])

h2(doc, "关于申请主体的说明")
bullet(doc, "以个人名义申请：材料简单，只需身份证复印件（正反面），无需单位盖章，审核较快；")
bullet(doc, "以学校名义申请：需提供事业单位法人证书复印件并加盖公章，通常由学校科研部门或"
            "信息中心统一办理；")
bullet(doc, "如为在校学生独立开发，建议以个人名义申请，权利归属清晰、办理更快。")

h1(doc, "六、需要一并准备的信息")
table(doc, [
    ["项目", "内容"],
    ["软件名称（务必一致）", f"{SOFTWARE} {VERSION}"],
    ["已上线访问地址", "https://xiaoaiyun.top（公开端）、https://admin.xiaoaiyun.top（管理端）"],
    ["部署服务器", "京东云（JD Cloud）云服务器，中国大陆地域"],
    ["源程序提交页数", "前 30 页 + 后 30 页，共 60 页"],
    ["文档提交页数", "《操作手册》全部 27 页"],
], [4.6, 11.4])

doc.save(OUT_DIR / "申请表填写信息.docx")
print("已生成: 申请表填写信息.docx")

# ==========================================================
# 文档二：材料清单与办理指引
# ==========================================================
doc2 = new_doc("材料清单与办理指引")
title1(doc2, "软件著作权登记 · 材料清单与办理指引")

h1(doc2, "一、需要提交的材料清单")
table(doc2, [
    ["序号", "材料名称", "份数/页数", "状态"],
    ["1", "软件著作权登记申请表", "在线填写并打印", "☐ 待填写"],
    ["2", "源程序（前 30 页 + 后 30 页）", "60 页", "☑ 已生成"],
    ["3", "软件文档（操作手册）", "27 页", "☑ 已生成"],
    ["4", "申请人身份证明", "身份证复印件 1 份", "☐ 待准备"],
    ["5", "其他证明文件（如有）", "按需", "—"],
], [1.6, 6.4, 4.2, 3.0])

h1(doc2, "二、材料准备要求")
h2(doc2, "2.1 源程序")
bullet(doc2, "提交源程序前、后各连续 30 页；", head="页数：")
bullet(doc2, "每页不少于 50 行（结束页除外）；", head="行数：")
bullet(doc2, f"页眉需标注软件名称及版本号（本文件已标注「{SOFTWARE} {VERSION} 源程序」）；", head="页眉：")
bullet(doc2, "页码连续编号（本文件已在页脚标注「第 N 页」）；", head="页码：")
bullet(doc2, "A4 纸单面打印，字迹清晰；", head="打印：")

h2(doc2, "2.2 文档（操作手册）")
bullet(doc2, "提交前 30 页 + 后 30 页，不足 60 页则提交全部；本手册共 27 页，全部提交；")
bullet(doc2, "页眉需标注软件名称及版本号；")
bullet(doc2, "包含软件功能、操作流程与界面截图，图文并茂。")

h2(doc2, "2.3 装订要求")
bullet(doc2, "源程序与文档分别装订，不要合订在一起；")
bullet(doc2, "建议在右上角简单标注「源程序」「文档」以便区分；")
bullet(doc2, "如在线申请，按要求上传 PDF 或图片格式即可，无需邮寄纸质件。")

h1(doc2, "三、办理流程")
step_text = [
    ("注册账号", "登录中国版权保护中心「软件著作权登记系统」（register.ccopyright.com.cn），"
                 "注册个人或单位账号并完成实名认证。"),
    ("在线填表", "按《申请表填写信息》文档逐项填写软件基本信息、开发环境、著作权人信息等。"),
    ("上传材料", "上传源程序文档、操作手册、身份证明扫描件。"),
    ("提交申请", "确认无误后提交，系统生成受理号和申请表 PDF。"),
    ("缴费", "按提示缴纳登记费（目前软件著作权登记已免收登记费，以官方页面为准）。"),
    ("受理审查", "版权中心受理并审查，一般 30–60 个工作日；"
                  "审查通过后颁发《计算机软件著作权登记证书》。"),
    ("领取证书", "可选择邮寄纸质证书或下载电子证书。"),
]
for idx, (name, desc) in enumerate(step_text, 1):
    p = doc2.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{idx}. {name}")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    para(doc2, desc, 10, GREY, indent=1.0)

h1(doc2, "四、注意事项")
bullet(doc2, "软件名称在申请表、源程序页眉、操作手册页眉中必须完全一致，"
             "否则会被要求补正；")
bullet(doc2, "源程序中如出现他人姓名、单位名称或第三方版权信息，建议先做处理；"
             "本项目的源程序已确认为自主开发代码；")
bullet(doc2, "首次发表日期可填软件首次上线日期；若尚未公开发表，则选择「未发表」；")
bullet(doc2, "版本号建议从 V1.0 开始，后续升级可另行登记；")
bullet(doc2, "登记完成后请妥善保管证书编号，可作为技术成果、评奖评优或保研加分的证明材料。")

h1(doc2, "五、本目录文件说明")
table(doc2, [
    ["文件名", "用途"],
    ["源程序-生活助手宿舍管理系统V1.0.docx", "源程序材料（60 页），直接打印或转 PDF 提交"],
    ["操作手册-生活助手宿舍管理系统V1.0.docx", "文档材料（27 页），直接打印或转 PDF 提交"],
    ["申请表填写信息.docx", "在线填表时对照填写"],
    ["材料清单与办理指引.docx", "本文件，办理流程参考"],
    ["screenshots/", "操作手册所用界面截图（16 张）"],
], [8.0, 8.0])

doc2.save(OUT_DIR / "材料清单与办理指引.docx")
print("已生成: 材料清单与办理指引.docx")
print()
print("输出目录:", OUT_DIR)
for f in sorted(OUT_DIR.iterdir()):
    if f.is_file():
        print(f"  {f.name:<46} {f.stat().st_size / 1024:>8.1f} KB")
