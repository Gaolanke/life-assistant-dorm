#!/usr/bin/env python3
"""生成《生活助手 · 班主任管理说明》Word 文档。"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = r"E:\Document\code\宿舍管理系统\docs\生活助手-班主任管理说明.docx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
GREY = RGBColor(0x55, 0x5F, 0x6D)

doc = Document()

# ---------- 全局字体 ----------
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.line_spacing = 1.4
style.paragraph_format.space_after = Pt(4)

USABLE_CM = 17.0  # A4 (21cm) 减去左右各 2cm 页边距

for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)


def set_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def heading(text, level=1):
    sizes = {0: 20, 1: 14, 2: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=sizes.get(level, 12), bold=True,
             color=ACCENT if level <= 1 else RGBColor(0x22, 0x2B, 0x38))
    return p


def para(text, bold=False, color=None, size=10.5, indent=0):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p


def bullet(text, bold_head=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_head:
        r1 = p.add_run(bold_head)
        set_font(r1, bold=True)
    r2 = p.add_run(text)
    set_font(r2)
    return p


def table(rows, widths=None, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(val))
            set_font(run, size=10, bold=(header and i == 0))
    if widths:
        total = sum(widths) or 1.0
        scaled = [w / total * USABLE_CM for w in widths]
        for row in t.rows:
            for idx, w in enumerate(scaled):
                row.cells[idx].width = Cm(w)
    return t


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    run = p.add_run("提示：" + text)
    set_font(run, size=9.5, color=GREY)
    return p


# ============================ 封面标题 ============================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
set_font(title.add_run("生活助手 · 班主任管理说明"), size=22, bold=True, color=ACCENT)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(14)
set_font(sub.add_run("兰州工业学院　电子信息工程学院　物联网 25-2 班"), size=11, color=GREY)

table([
    ["项目", "内容"],
    ["公开填报网址", "https://xiaoaiyun.top"],
    ["管理端网址", "https://admin.xiaoaiyun.top"],
    ["管理账号", "admin"],
    ["管理密码", "344940（请登录后自行修改）"],
    ["夜间填报截止", "每天 22:00"],
    ["群内自动汇报", "每天 22:40"],
    ["宿舍 / 学生", "14 间宿舍　44 名同学"],
], widths=[4.0, 10.5])

# ============================ 一、系统简介 ============================
heading("一、这个系统是做什么的", 1)
para("「生活助手」是本班生活委员搭建的宿舍管理网页系统，用于替代每天在群里逐间询问、手工统计的填报方式。")
para("它解决三件事：")
bullet("同学不用注册账号，扫码或点链接就能填，点一下自己的宿舍即可；", "① 填报方便：")
bullet("系统自动统计哪间填了、哪间没填、谁缺席、原因是什么；", "② 自动统计：")
bullet("每天 22:40 自动把当晚汇总发到班级 QQ 群，老师打开手机就能看到。", "③ 自动汇报：")

# ============================ 二、访问方式 ============================
heading("二、怎么访问", 1)
para("管理端（老师用）", bold=True, color=ACCENT)
para("网址：https://admin.xiaoaiyun.top", indent=0.4)
para("登录：账号 admin，密码 344940", indent=0.4)
note("首次登录后建议点左下角「修改密码」改掉初始密码。管理端网址请勿转发给学生。")

para("填报端（学生用）", bold=True, color=ACCENT)
para("网址：https://xiaoaiyun.top", indent=0.4)
para("学生不需要账号密码，进入后点击自己所在的宿舍即可开始填写。", indent=0.4)

# ============================ 三、功能说明 ============================
heading("三、管理端有哪些功能", 1)

heading("1. 数据看板", 2)
para("打开就能看到今晚的填报进度：多少间已填、多少间未填、在寝人数、缺席人次，以及未填报宿舍的名单。")

heading("2. 统计图表", 2)
para("按「日 / 周 / 月」三个维度查看历史趋势，包括填报率变化、缺席情况分布、各宿舍表现对比等，图表直观可见。")

heading("3. 填报记录", 2)
para("每一间宿舍每天的填报都在这里，可按日期查看详情（谁在、谁不在、原因、备注）。")
para("支持一键导出 Excel：系统每天 23:30 还会自动存一份到服务器，长期留档不会丢。", indent=0.4)

heading("4. 填写日志", 2)
para("记录每一次填报操作的时间、宿舍、设备来源，便于核查（例如出现争议时可追溯）。")

heading("5. 宿舍与人员", 2)
para("管理全部宿舍和人员名单，信息包含：姓名、学号、电话、宿舍号。")
bullet("支持 Excel 批量导入名单；")
bullet("可自由增加、删除、修改任何一条人员信息。")

heading("6. 班费管理", 2)
para("记录班费收支，数据来源于导入的班费收支明细表，可随时修改编辑。")
table([
    ["项目", "金额"],
    ["班费总额", "2400.00 元（48 人 × 50 元）"],
    ["已使用", "497.85 元"],
    ["当前余额", "1902.15 元"],
], widths=[6.0, 8.5])
para("支持把最新班费情况一键发到班级群公示。", indent=0.4)

heading("7. 提醒推送", 2)
para("查看每晚自动汇报的实际内容，也可以在这里手动补发一次。")
para("推送记录全部留痕，成功还是失败都能查到。", indent=0.4)

heading("8. 账号管理", 2)
para("为每间宿舍的寝室长设置登录账号（账号就是宿舍号，如 4-153）。")
bullet("不设密码：任何人都能帮该宿舍填报；")
bullet("设了密码：只有该寝室长能填，其他人无法修改。")
note("管理员账号不受限制，任何时候都能修改任何宿舍的数据。")

heading("9. 系统配置", 2)
para("所有参数都可以在网页上图形化修改，不需要改代码：学校名称、班级、截止时间、汇报时间、缺席原因选项、QQ 官方机器人、Excel 自动导出等。")

# ============================ 四、常用操作 ============================
doc.add_page_break()
heading("四、四个最常用的操作", 1)

heading("操作 1：查看今晚谁没填", 2)
para("登录管理端 →「数据看板」→ 页面右侧「还差这些寝室没填」即为未填报名单。", indent=0.4)

heading("操作 2：导出某天的填报记录", 2)
para("登录管理端 →「填报记录」→ 选择日期 → 点「导出 Excel」。", indent=0.4)

heading("操作 3：修改一间宿舍的人员名单", 2)
para("登录管理端 →「宿舍与人员」→ 找到对应宿舍 → 修改后保存即可。", indent=0.4)

heading("操作 4：修改自己的密码", 2)
para("登录管理端 → 页面左下角「修改密码」→ 输入旧密码和新密码。", indent=0.4)

# ============================ 五、自动任务 ============================
heading("五、系统自动做的事（不用手动操作）", 1)
table([
    ["时间", "自动执行的事"],
    ["每天 22:00", "填报截止，之后系统开始统计未填报宿舍"],
    ["每天 22:40", "自动把当晚填报汇总发到班级 QQ 群"],
    ["每天 23:30", "自动导出当天 Excel 并存到服务器留档"],
], widths=[3.5, 11.0])
note("如遇网络波动，群消息可能延迟几分钟；系统会自动重试。")

# ============================ 六、常见问题 ============================
heading("六、常见问题", 1)

qa = [
    ("同学说打不开网页怎么办？",
     "先确认网址是 https://xiaoaiyun.top（要带 https）。若家里网络异常，可改用备用地址 http://111.228.38.170:8090 访问。"),
    ("同学填错了能改吗？",
     "可以。同一宿舍当天重新填写一次即可覆盖；管理员也能直接在「填报记录」里修改。"),
    ("群里没收到 22:40 的汇报？",
     "可在管理端「提醒推送」页面点「立即发送到 QQ 群」手动补发；页面上能看到每次发送的成功或失败记录。"),
    ("忘记管理密码了？",
     "可联系生活委员在服务器上重置，无需重装系统。"),
    ("数据会不会丢？",
     "所有数据保存在服务器上，且每天自动导出 Excel 备份。管理端「系统与备份」页面还可以随时下载一份完整数据库备份。"),
]
for q, a in qa:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    set_font(p.add_run("问：" + q), bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(0.4)
    set_font(p2.add_run("答：" + a))

# ============================ 七、安全提示 ============================
heading("七、安全提示", 1)
bullet("管理端网址和密码请勿发到学生群，避免无关人员登录；")
bullet("建议登录后立即修改初始密码 344940；")
bullet("如需更高安全性，可为管理端开启「邮箱验证码」双重保护（可让生活委员配置）；")
bullet("本系统仅供班级内部使用，请勿对外公开传播。")

# ============================ 结尾 ============================
end = doc.add_paragraph()
end.paragraph_format.space_before = Pt(18)
set_font(end.add_run("技术问题请联系本班生活委员。"), size=10, color=GREY)

doc.save(OUT)
print("已生成：", OUT)
print("大小：", os.path.getsize(OUT), "字节")
