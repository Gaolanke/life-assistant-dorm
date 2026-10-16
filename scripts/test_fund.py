"""班费模块端到端测试（全部走 HTTP 接口，用真实的班费收支明细表）。

覆盖：
  上传真实 xlsx 预览 → 导入入库 → 概览校验（每人 50 / 人数 48 / 合计 2400 /
  已用 497.85 / 余额 1902.15）→ 明细增删改 → 发群文案预览 → 导出 xlsx → 权限。

用法： .venv\\Scripts\\python.exe scripts\\test_fund.py [管理端口] [公开端口]
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
import server_env  # noqa: E402

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout.reconfigure(encoding="utf-8")

from auth_client import (  # noqa: E402
    ADMIN_BASE,
    PUBLIC_BASE,
    ensure_admin_client,
    public_client,
)

ADMIN = sys.argv[1] if len(sys.argv) > 1 else ADMIN_BASE
PUBLIC = sys.argv[2] if len(sys.argv) > 2 else PUBLIC_BASE

def _find_xlsx() -> Path:
    """找班费明细表：优先环境变量，其次项目根目录，最后原始路径。"""
    import os

    candidates = []
    env = os.environ.get("DORM_FUND_XLSX", "").strip()
    if env:
        candidates.append(Path(env))
    candidates.append(ROOT / "班费收支明细表 (12).xlsx")
    candidates.append(ROOT / "data" / "班费收支明细表 (12).xlsx")
    candidates.append(
        Path(
            r"C:\Users\Xiaoai\Documents\xwechat_files\wxid_l49z2izypix422_7850"
            r"\msg\file\2026-10\班费收支明细表 (12).xlsx"
        )
    )
    for c in candidates:
        if c.exists():
            return c
    return candidates[0]


XLSX = _find_xlsx()

# 真实表格里写死的表头统计
EXPECT_PER_PERSON = 50.0
EXPECT_HEADCOUNT = 48
EXPECT_TOTAL = 2400.0
EXPECT_SPENT = 497.85
EXPECT_BALANCE = 1902.15

PASS, FAIL = [], []


def check(name: str, ok: bool, detail: str = "") -> None:
    (PASS if ok else FAIL).append(name if ok else f"{name} :: {detail}")
    print(f"  [{'OK' if ok else 'FAIL'}] {name}{'' if ok else '  ' + detail}")


def close(a: float, b: float) -> bool:
    return abs(float(a) - float(b)) < 0.01


def main() -> int:
    print(f"数据目录：{server_env.DATA_DIR}")
    print(f"测试表格：{XLSX}")
    if not XLSX.exists():
        print(f"  找不到真实班费表：{XLSX}")
        return 2
    content = XLSX.read_bytes()

    admin = ensure_admin_client(base=ADMIN)
    check("管理员登录成功", bool(admin.headers.get("X-Session-Token")), "没有拿到会话令牌")

    print("\n== 1. 上传真实表格预览 ==")
    preview_res = admin.post(
        "/api/fund/import/preview",
        files={
            "file": (
                XLSX.name,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    check("预览接口返回 200", preview_res.status_code == 200, preview_res.text[:200])
    preview = preview_res.json()
    check("预览 ok=true", preview.get("ok") is True, str(preview)[:200])
    entry_count = int(preview.get("entry_count") or 0)
    check("解析出明细行", entry_count >= 1, f"entry_count={entry_count}")
    check("识别到支出项目与金额列", "item" in preview.get("columns", {}) and "amount" in preview.get("columns", {}), str(preview.get("columns")))
    check("明细合计与表里「已使用」一致", preview.get("spent_match") is True, f"parsed={preview.get('spent_parsed')} sheet={preview.get('spent_in_sheet')}")
    summary = preview.get("summary") or {}
    check("表头每人=50", close(summary.get("per_person", 0), EXPECT_PER_PERSON), str(summary))
    check("表头人数=48", int(summary.get("headcount") or 0) == EXPECT_HEADCOUNT, str(summary))
    check("表头班费合计=2400", close(summary.get("total_fund", 0), EXPECT_TOTAL), str(summary))
    check("表头已使用=497.85", close(summary.get("spent", 0), EXPECT_SPENT), str(summary))
    check("表头余额=1902.15", close(summary.get("balance", 0), EXPECT_BALANCE), str(summary))
    first = (preview.get("entries") or [{}])[0]
    check("明细行带项目名与金额", bool(first.get("item")) and float(first.get("amount") or 0) > 0, str(first))

    print("\n== 2. 导入入库并校验概览 ==")
    commit = admin.post("/api/fund/import/commit", json={
        "entries": preview["entries"],
        "summary": preview["summary"],
        "replace": True,
        "apply_summary": True,
    })
    check("导入提交成功", commit.status_code == 200 and commit.json().get("ok"), commit.text[:200])
    check("导入条数与预览一致", int(commit.json().get("created") or 0) == entry_count, commit.text[:200])

    summary_res = admin.get("/api/fund/summary")
    check("概览接口可用", summary_res.status_code == 200, str(summary_res.status_code))
    stats = summary_res.json()
    check("概览：每人 50", close(stats.get("per_person"), EXPECT_PER_PERSON), str(stats))
    check("概览：人数 48", int(stats.get("headcount") or 0) == EXPECT_HEADCOUNT, str(stats))
    check("概览：班费合计 2400", close(stats.get("total_fund"), EXPECT_TOTAL), str(stats))
    check("概览：已使用 497.85", close(stats.get("spent"), EXPECT_SPENT), str(stats))
    check("概览：余额 1902.15", close(stats.get("balance"), EXPECT_BALANCE), str(stats))
    check("概览：明细条数正确", int(stats.get("entry_count") or 0) == entry_count, str(stats))

    entries_res = admin.get("/api/fund/entries")
    entries = entries_res.json()
    check("明细列表可用", entries_res.status_code == 200 and entries.get("total") == entry_count, str(entries)[:150])
    check("明细列表带概览", close((entries.get("summary") or {}).get("spent"), EXPECT_SPENT), str(entries.get("summary")))
    check("明细按日期排序且带序号", all(item.get("seq") for item in entries["items"]), str(entries["items"][:1]))

    saved = admin.put("/api/fund/settings", json={"per_person": 50, "headcount": 48, "class_name": "物联网25-2"})
    check("班费设置可保存", saved.status_code == 200 and saved.json()["ok"], saved.text[:200])
    check("保存后概览跟着变", saved.json()["summary"]["class_name"] == "物联网25-2", saved.text[:200])

    print("\n== 3. 明细增删改 ==")
    created = admin.post("/api/fund/entries", json={
        "date": "2026-10-10", "item": "自检：班费打印费", "amount": 12.5,
        "handler": "生活委员", "checker": "杨朝辉", "note": "自动测试",
    })
    check("新增明细成功", created.status_code == 200 and created.json()["ok"], created.text[:200])
    entry = created.json()["entry"]
    check("新明细自动分配序号", int(entry.get("seq") or 0) > 0, str(entry))
    check("新增后已使用 +12.5", close(created.json()["summary"]["spent"], EXPECT_SPENT + 12.5), str(created.json()["summary"]))

    edited = admin.patch(f"/api/fund/entries/{entry['id']}", json={"amount": 20, "note": "改成20"})
    check("修改明细成功", edited.status_code == 200 and close(edited.json()["entry"]["amount"], 20), edited.text[:200])
    check("修改后已使用 +20", close(edited.json()["summary"]["spent"], EXPECT_SPENT + 20), str(edited.json()["summary"]))

    found = admin.get("/api/fund/entries?keyword=班费打印费").json()
    check("按关键字搜索到新明细", found["total"] == 1 and found["items"][0]["id"] == entry["id"], str(found)[:150])

    missing = admin.patch("/api/fund/entries/99999999", json={"amount": 1})
    check("改不存在的明细返回 404", missing.status_code == 404, str(missing.status_code))

    empty_delete = admin.post("/api/fund/entries/delete", json={"ids": []})
    check("不选记录删除会被拒绝", empty_delete.status_code == 400, str(empty_delete.status_code))

    deleted = admin.post("/api/fund/entries/delete", json={"ids": [entry["id"]]})
    check("删除明细成功", deleted.status_code == 200 and deleted.json()["deleted"] == 1, deleted.text[:200])
    check("删除后已使用回到 497.85", close(deleted.json()["summary"]["spent"], EXPECT_SPENT), str(deleted.json()["summary"]))

    print("\n== 4. 发群文案预览 ==")
    latest = admin.post("/api/fund/push/preview", json={"scope": "latest"}).json()
    check("最新一笔预览可用", latest.get("ok") is True and latest.get("count") == 1, str(latest)[:200])
    check("文案含「班费」标题", "班费" in (latest.get("text") or ""), (latest.get("text") or "")[:120])
    check("文案含每人/人数/合计", "50" in latest["text"] and "48" in latest["text"] and "2400" in latest["text"], latest["text"][:200])
    check("文案含已使用与余额", "497.85" in latest["text"] and "1902.15" in latest["text"], latest["text"][:200])
    check("文案含明细行", "元（经手人：" in latest["text"], latest["text"][:200])

    all_res = admin.post("/api/fund/push/preview", json={"scope": "all"}).json()
    check("全部明细预览条数正确", all_res["count"] == entry_count, str(all_res)[:150])
    recent = admin.post("/api/fund/push/preview", json={"scope": "recent", "limit": 3}).json()
    check("最近 3 笔预览条数正确", recent["count"] == 3, str(recent)[:150])
    ids = [item["id"] for item in entries["items"][:2]]
    selected = admin.post("/api/fund/push/preview", json={"scope": "selected", "ids": ids}).json()
    check("选定明细预览条数正确", selected["count"] == 2, str(selected)[:150])
    no_selection = admin.post("/api/fund/push/preview", json={"scope": "selected", "ids": []}).json()
    check("没勾选时给出提示", "勾选" in (no_selection.get("text") or ""), str(no_selection)[:150])

    print("\n== 5. 导出 xlsx ==")
    export = admin.get("/api/fund/export")
    check("导出接口返回 xlsx", export.status_code == 200 and export.content[:2] == b"PK", str(export.status_code))
    wb = load_workbook(io.BytesIO(export.content))
    ws = wb[wb.sheetnames[0]]
    check("导出的工作表叫「班费明细」", ws.title == "班费明细", ws.title)
    header = [cell.value for cell in ws[4]]
    check("导出表头正确", "支出项目" in header and "金额" in header and "经手人" in header, str(header))
    check("导出含表头统计", close(ws["E2"].value, EXPECT_PER_PERSON) and int(ws["F2"].value or 0) == EXPECT_HEADCOUNT, f"E2={ws['E2'].value} F2={ws['F2'].value}")
    check("导出明细条数正确", (ws.max_row - 4) == entry_count, f"max_row={ws.max_row}")

    print("\n== 6. 权限 ==")
    with httpx.Client(base_url=ADMIN, timeout=40, follow_redirects=True) as anon:
        check("未登录看不了班费概览", anon.get("/api/fund/summary").status_code == 401, "应返回 401")
        check("未登录看不了班费明细", anon.get("/api/fund/entries").status_code == 401, "应返回 401")
        check("未登录导不了班费表", anon.get("/api/fund/export").status_code == 401, "应返回 401")
        check("未登录改不了班费设置", anon.put("/api/fund/settings", json={"per_person": 1}).status_code == 401, "应返回 401")
    public = public_client(PUBLIC)
    check("公开端口没有班费接口", public.get("/api/fund/summary").status_code == 404, "公开端口不应该有班费接口")

    admin.close()
    public.close()

    print("\n" + "=" * 60)
    print(f"通过 {len(PASS)} 项，失败 {len(FAIL)} 项")
    for item in FAIL:
        print("  - " + item)
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
