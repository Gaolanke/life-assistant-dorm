#!/usr/bin/env python3
"""生成《生活助手 · 使用说明》（含同学填报篇 + 管理端篇）。"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = r"E:\Document\code\宿舍管理系统\docs\生活助手-使用说明.docx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
GREY = RGBColor(0x5A, 0x63, 0x72)
DARK = RGBColor(0x1F, 0x29, 0x37)

doc = Document()

style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.line_spacing = 1.45
style.paragraph_format.space_after = Pt(4)

USABLE_CM = 17.0  # A4 (21cm) 减去左右各 2cm 页边距

for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)


def set_font(run, size=None, bold=None, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def heading(text, level=1):
    sizes = {0: 21, 1: 14.5, 2: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level <= 1 else 11)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), size=sizes.get(level, 12), bold=True,
             color=ACCENT if level <= 1 else DARK)
    return p


def para(text, bold=False, color=None, size=10.5, indent=0, space_before=0):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if space_before:
        p.paragraph_format.space_before = Pt(space_before)
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def bullet(text, head=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if head:
        set_font(p.add_run(head), bold=True)
    set_font(p.add_run(text))
    return p


def step(num, text, detail=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run(f"{num}. "), bold=True, color=ACCENT)
    set_font(p.add_run(text))
    if detail:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(1.2)
        p2.paragraph_format.space_after = Pt(5)
        set_font(p2.add_run(detail), size=9.5, color=GREY)


def table(rows, widths=None, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run(str(val)), size=10, bold=(header and i == 0))
    if widths:
        total = sum(widths) or 1.0
        scaled = [w / total * USABLE_CM for w in widths]
        for row in t.rows:
            for idx, w in enumerate(scaled):
                row.cells[idx].width = Cm(w)
    return t


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(3)
    set_font(p.add_run("提示　"), size=9.5, bold=True, color=ACCENT)
    set_font(p.add_run(text), size=9.5, color=GREY)
    return p


# ======================= 封面 =======================
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_after = Pt(2)
set_font(t.add_run("生活助手 · 使用说明"), size=23, bold=True, color=ACCENT)

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.paragraph_format.space_after = Pt(16)
set_font(s.add_run("兰州工业学院　电子信息工程学院　物联网 25-2 班"), size=11, color=GREY)

table([
    ["用途", "网址 / 账号"],
    ["同学、寝室长填报", "https://xiaoaiyun.top"],
    ["管理员（生活委员 / 老师）", "https://admin.xiaoaiyun.top"],
    ["管理账号", "admin"],
    ["管理密码", "344940"],
], widths=[5.0, 9.5])

note("本说明分两部分：第一部分给同学和寝室长看，第二部分给管理端使用者看。")

# ======================= 第一部分 =======================
doc.add_page_break()
h = doc.add_paragraph()
set_font(h.add_run("第一部分　同学 / 寝室长：怎么填报"), size=17, bold=True, color=ACCENT)
para("")

heading("一、打开填报页面", 1)
para("在手机或电脑浏览器里输入网址：", bold=False)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(0.8)
set_font(p.add_run("https://xiaoaiyun.top"), size=13, bold=True, color=ACCENT)
note("网址开头是 https（有一个小锁图标），不要漏掉 s。建议把这个网址收藏到浏览器，或用手机扫码后加书签。")

heading("二、填写步骤（约 20 秒）", 1)
step(1, "点一下自己所在的宿舍卡片", "页面上会列出全部宿舍，例如「6-213　杨朝辉·4 人」。也可以用上面的搜索框输入宿舍号快速查找。")
step(2, "在弹出/跳转的页面里确认人员是否齐全", "系统会列出该宿舍的全部成员，逐个勾选在寝或不在。")
step(3, "如果有人在寝之外的情况，选择原因并填写备注", "原因可选：回家、病假、事假、外出、实习、其他。")
step(4, "点「提交」，看到成功提示即完成", "当天可以重复填写，重新提交会覆盖之前的内容。")

heading("三、填写时间要求", 1)
table([
    ["时间", "要求"],
    ["每天 22:00 前", "请各寝室完成填报"],
    ["每天 22:40", "系统自动把汇总发到班级 QQ 群，老师查看"],
], widths=[4.5, 10.0])

heading("四、寝室长说明", 1)
para("如果本宿舍设置了密码，只有寝室长能填报；没有设置密码的宿舍，任何同学都可以代为填报。")
bullet("寝室长的登录账号就是宿舍号，例如 6-213；")
bullet("密码由生活委员设置，忘记了可以找生活委员重置；")
bullet("登录入口在填报页面首页的「我是寝室长（验证）」按钮。")

heading("五、填报遇到问题", 1)
bullet("打不开网页：确认网址是 https://xiaoaiyun.top，换个网络或改用 http://111.228.38.170:8090 试试；")
bullet("找不到自己的宿舍：联系生活委员核实宿舍号；")
bullet("填错了：当天重新填一次即可覆盖；")
bullet("宿舍号/人员名单有误：联系生活委员修改。")

# ======================= 第二部分 =======================
doc.add_page_break()
h2 = doc.add_paragraph()
set_font(h2.add_run("第二部分　管理端：怎么查看和管理"), size=17, bold=True, color=ACCENT)
para("")

heading("一、登录管理端", 1)
step(1, "打开 https://admin.xiaoaiyun.top")
step(2, "输入账号 admin，密码 344940")
step(3, "勾选「记住我」可在 30 天内免登录")
note("管理端网址和密码请勿发到学生群。首次登录后建议点左下角「修改密码」改掉初始密码。")

heading("二、功能一览", 1)
table([
    ["页面", "用途"],
    ["数据看板", "今晚填报进度：已填/未填宿舍、在寝人数、缺席人次、未填报名单"],
    ["统计图表", "按日 / 周 / 月查看填报率、缺席情况等趋势图"],
    ["填报记录", "逐条查看每天每间宿舍的填报详情，可导出 Excel"],
    ["填写日志", "每次填报的时间、宿舍、设备留痕，便于核查"],
    ["班费管理", "班费收支明细，可编辑、可导入导出、可一键发到班级群"],
    ["宿舍与人员", "管理宿舍和名单（姓名、学号、电话、宿舍号），支持 Excel 批量导入"],
    ["提醒推送", "查看/手动补发每晚群汇报与未填报提醒，含发送成功失败记录"],
    ["账号管理", "为寝室长设置或重置登录密码（账号＝宿舍号）"],
    ["管理控制台", "宿舍长密码总览：批量设置、重置、生成账号"],
    ["系统配置", "学校名称、截止时间、汇报时间、缺席原因、QQ 官方机器人等"],
], widths=[3.6, 11.0])

heading("三、四个最常用的操作", 1)

para("操作 1　查看今晚谁没填", bold=True, color=DARK, space_before=6)
para("「数据看板」→ 页面右下角「还差这些寝室没填」即为未填报名单。", indent=0.5)

para("操作 2　导出某天的填报记录", bold=True, color=DARK, space_before=6)
para("「填报记录」→ 选择日期 → 点「导出 Excel」。", indent=0.5)

para("操作 3　修改一间宿舍的人员名单", bold=True, color=DARK, space_before=6)
para("「宿舍与人员」→ 找到该宿舍 → 修改后保存。", indent=0.5)

para("操作 4　把班费发到班级群", bold=True, color=DARK, space_before=6)
para("「班费管理」→「发到班级群」→ 选择发送范围（最新一笔 / 最近几笔 / 选定 / 全部）→ 先「预览要发的内容」→ 确认后「直接发送到群」。", indent=0.5)

heading("四、系统自动做的事", 1)
table([
    ["时间", "自动执行"],
    ["每天 22:00", "填报截止，开始统计未填报宿舍"],
    ["每天 22:40", "自动把当晚填报汇总发到班级 QQ 群"],
    ["每天 23:30", "自动导出当天 Excel 并存到服务器留档"],
], widths=[3.5, 11.0])
note("群消息由 QQ 官方机器人发送；如果网络波动导致发送失败，系统会自动重试。")

heading("五、常见问题", 1)
qa = [
    ("群里没收到 22:40 的汇报怎么办？",
     "打开「提醒推送」页面 → 点「立即发送到 QQ 群」手动补发；页面下方能看到每次发送的成功或失败原因。"),
    ("同学填错了能改吗？",
     "能。同一宿舍当天重新提交会覆盖；管理员也可在「填报记录」里直接修改。"),
    ("怎么给寝室长设密码？",
     "「账号管理」里找到该宿舍 → 输入密码保存。也可以「管理控制台」里一键批量设置。设了密码后，只有该寝室长能填报。"),
    ("忘记管理密码了怎么办？",
     "联系生活委员在服务器上重置，不用重装系统。"),
    ("数据会不会丢？",
     "数据全部存在服务器上，且每天自动导出 Excel 备份；「系统与备份」页面还可以随时下载完整数据库备份。"),
    ("能导出全班汇总表吗？",
     "可以。「填报记录」页面支持导出 Excel，包含每间宿舍每天的填报详情。"),
]
for q, a in qa:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    set_font(p.add_run("问　" + q), bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(0.5)
    set_font(p2.add_run("答　" + a))

heading("六、安全提示", 1)
bullet("管理端网址和密码不要发到学生群；")
bullet("登录后建议立即修改初始密码 344940；")
bullet("本系统仅限班级内部使用，请勿对外公开传播；")
bullet("如需更高安全性，可为管理端加一层邮箱验证码保护（可让生活委员配置）。")

end = doc.add_paragraph()
end.paragraph_format.space_before = Pt(20)
set_font(end.add_run("技术问题请联系本班生活委员。"), size=10, color=GREY)

doc.save(OUT)
print("已生成：", OUT)
print("大小：", os.path.getsize(OUT), "字节")
